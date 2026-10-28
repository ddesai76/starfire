#!/usr/bin/env python3

# starfire.py:    TEST CAMPAIGN / TEST CARD DESIGNER
# AUTHOR:         DANIEL DESAI
# UPDATED:        2026-10-27
# VERSION:        4.0.1
#

"""
Usage
-----
    python3 starfire.py                    # Open GUI on http://localhost:5791
    python3 starfire.py --port 8080
    python3 starfire.py --no-browser

"""

from __future__ import annotations

import argparse
import base64
import binascii
import csv
import errno
import io
import json
import os
import re
import shlex
import shutil
import signal
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.request
import webbrowser
from dataclasses import dataclass, field, fields, asdict
from datetime import date, datetime, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs

try:
    import signet   # the Signet language and .sig files; optional, built into the binary when present
except ImportError:
    signet = None
if signet is not None and getattr(signet, "API", None) != 2:
    print(f"starfire: ignoring {signet.__file__} (API {getattr(signet, 'API', None)}, needs 2); Signals are off", file=sys.stderr)
    signet = None
SIGNET_AVAILABLE = signet is not None   # without it: no Signals box, live list, System Reports or stamping
_NO_SIGNET = "Signet isn't available (no signet.py next to starfire.py when it was built), so Signals are off"

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# Test card data model
@dataclass
class TestPoint:
    seq: int
    action: str = ""
    expected: str = ""
    observed: str = ""
    reference: str = ""
    responsible_party: str = ""
    date: str = ""
    time: str = ""        # 24h "HH:MM" -- filled by the UI's "now" button
    initial: str = ""
    quality: str = ""
    locked: bool = False
    flagged: bool = False
    images: list = field(default_factory=list)
    files: list = field(default_factory=list)      # uploaded with the + button
    parameters: list = field(default_factory=list)  # [[name, value]], passed to modules
    signals: str = ""                                # Signet code, run before a module
    runs: dict = field(default_factory=dict)         # module stem -> last run status
    module: str = ""                                 # the module the point runs, and its arguments
    args: str = ""


def _normalize_pair_list(items: list, width: int) -> list:
    """Shape normalization for nomenclature/references/
    notes/warnings: pads each row to `width` elements, and
    accepts a bare string in place of a single-element row."""
    out = []
    for item in items:
        row = [item] if isinstance(item, str) else list(item)
        while len(row) < width - 1:
            row.append("")
        if len(row) < width:
            row.append(False)
        out.append(row[:width])
    return out


@dataclass
class TestCard:
    id: str = "TC-001"
    title: str = "Untitled Test Card"
    revision: str = "A"
    author: str = ""
    date: str = field(default_factory=lambda: date.today().isoformat())
    project: str = ""
    test_conductor: str = ""
    quality_assurance: str = ""
    system: str = ""
    subsystem: str = ""
    test_type: str = ""
    article: str = ""
    requirements: str = ""
    points: list = field(default_factory=list)         # list[TestPoint]
    nomenclature: list = field(default_factory=list)
    references: list = field(default_factory=list)
    notes: list = field(default_factory=list)       # list[[text, locked]]
    warnings: list = field(default_factory=list)     # list[[text, locked]]
    locked: bool = False
    files_migrated: bool = False   # set once files already named for a point were attached

    @staticmethod
    def from_dict(d: dict) -> "TestCard":
        # Unknown keys are dropped, so files saved by an older version still open.
        known = {f.name for f in fields(TestPoint)}
        pts = [TestPoint(**{k: v for k, v in p.items() if k in known}) for p in d.get("points", [])]
        nom = [tuple(x) for x in _normalize_pair_list(d.get("nomenclature", []), 4)]
        refs = _normalize_pair_list(d.get("references", []), 3)
        notes = _normalize_pair_list(d.get("notes", []), 2)
        warnings = _normalize_pair_list(d.get("warnings", []), 2)
        base = {k: v for k, v in d.items()
                if k not in ("points", "nomenclature", "references", "notes", "warnings")}
        return TestCard(points=pts, nomenclature=nom, references=refs,
                         notes=notes, warnings=warnings, **base)

    def to_dict(self) -> dict:
        d = asdict(self)
        return d


def sorted_nomenclature(card: TestCard) -> list:
    # card.nomenclature as a list of (symbol, definition, units, locked),alphabetized.
    rows = [(r[0], r[1], r[2] if len(r) > 2 else "", r[3] if len(r) > 3 else False)
            for r in card.nomenclature]
    return sorted(rows, key=lambda row: row[0].lower())


# LaTeX Greek-letter commands -> Unicode, for the print/HTML view.
GREEK_UNICODE = {
    "alpha": "\u03b1", "beta": "\u03b2", "gamma": "\u03b3", "delta": "\u03b4",
    "epsilon": "\u03b5", "zeta": "\u03b6", "eta": "\u03b7", "theta": "\u03b8",
    "iota": "\u03b9", "kappa": "\u03ba", "lambda": "\u03bb", "mu": "\u03bc",
    "nu": "\u03bd", "xi": "\u03be", "pi": "\u03c0", "rho": "\u03c1",
    "sigma": "\u03c3", "tau": "\u03c4", "upsilon": "\u03c5", "phi": "\u03c6",
    "chi": "\u03c7", "psi": "\u03c8", "omega": "\u03c9",
    "Gamma": "\u0393", "Delta": "\u0394", "Theta": "\u0398", "Lambda": "\u039b",
    "Xi": "\u039e", "Pi": "\u03a0", "Sigma": "\u03a3", "Upsilon": "\u03a5",
    "Phi": "\u03a6", "Psi": "\u03a8", "Omega": "\u03a9",
}
_GREEK_CMD_RE = re.compile(r"\\([A-Za-z]+)")


def symbol_display(sym: str) -> str:
    """Render a nomenclature symbol for CSV/HTML: LaTeX Greek commands
    become their Unicode letter; anything else passes unchanged."""
    return _GREEK_CMD_RE.sub(lambda m: GREEK_UNICODE.get(m.group(1), m.group(0)), sym or "")


# Export as CSV
def export_csv(card: TestCard) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["ID", "Title", "Revision", "Project", "Test Conductor",
                "Quality Assurance", "Author", "Date", "System", "Subsystem", "Test Type",
                "Requirement(s)"])
    w.writerow([card.id, card.title, card.revision, card.project, card.test_conductor,
                card.quality_assurance, card.author, card.date, card.system,
                card.subsystem, card.test_type, card.requirements])
    w.writerow([])
    w.writerow(["Seq", "Action", "Expected Result", "Observed Result", "Reference",
                "Responsible Party", "Date", "Time", "Initial", "Quality"])
    for p in card.points:
        w.writerow([p.seq, p.action, p.expected, p.observed, p.reference,
                    p.responsible_party, p.date, p.time, p.initial, p.quality])
    return buf.getvalue()


# Export as print-ready HTML (browser print-to-PDF, no new dependency).
def _h(s: str) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def _point_params(p: TestPoint) -> tuple:
    # (names, values), one per line, for a card's Parameter and Value columns.
    rows = [(str(r[0]).strip(), str(r[1])) for r in (p.parameters or []) if len(r) > 1 and str(r[0]).strip()]
    return "\n".join(n for n, _ in rows), "\n".join(v for _, v in rows)


def _params_line(p: TestPoint) -> str:
    # "voltage = 5, current = 1" for a Test Report; "" with no parameters.
    return ", ".join(f"{str(r[0]).strip()} = {r[1]}" for r in (p.parameters or []) if len(r) > 1 and str(r[0]).strip())


def export_print_html(card: TestCard, host: str = "localhost") -> str:
    # Parameter and Value columns only when some point has parameters.
    has_params = any(_point_params(p)[0] for p in card.points)
    def pcells(p):
        if not has_params:
            return ""
        names, values = _point_params(p)
        return f"<td>{_h(names).replace(chr(10), '<br>')}</td><td>{_h(values).replace(chr(10), '<br>')}</td>"
    rows = "\n".join(
        f"<tr><td>{p.seq}</td><td>{_h(p.action)}</td>{pcells(p)}<td>{_h(p.expected)}</td>"
        f"<td>{_h(p.observed)}</td><td>{_h(p.reference)}</td>"
        f"<td>{_h(p.responsible_party)}</td><td>{_h(p.date)}</td><td>{_h(p.time)}</td>"
        f"<td>{_h(p.initial)}</td><td>{_h(p.quality)}</td></tr>"
        for p in card.points
    )
    nom = sorted_nomenclature(card)
    nom_rows = "\n".join(
        f"<tr><td>{_h(symbol_display(sym))}</td><td>{_h(defn)}</td><td>{_h(units)}</td></tr>"
        for sym, defn, units, _locked in nom
    )
    nom_section = "" if not nom else f"""<h2>Nomenclature</h2>
<table><tr><th style="width:15%">Symbol</th><th>Definition</th><th style="width:15%">Units</th></tr>
{nom_rows}
</table>
"""
    refs = [(r, d) for r, d, _locked in card.references if (r or "").strip() or (d or "").strip()]
    ref_rows = "\n".join(f"<tr><td>{_h(r)}</td><td>{_h(d)}</td></tr>" for r, d in refs)
    ref_section = "" if not refs else f"""<h2>References</h2>
<table><tr><th style="width:20%">Reference</th><th>Description</th></tr>
{ref_rows}
</table>
"""
    warnings = [w for w, _locked in card.warnings if w.strip()]
    warn_bars = "\n".join(f'<div class="warn-bar">\u26a0 {_h(w)}</div>' for w in warnings)
    warn_section = "" if not warnings else f"""<h2>Warnings</h2>
{warn_bars}
"""
    widths = ([4, 14, 8, 6, 14, 14, 8, 7, 7, 5, 5, 8] if has_params else [4, 16, 16, 16, 10, 8, 8, 6, 6, 10])
    cols = "<colgroup>" + "".join(f'<col style="width:{w}%">' for w in widths) + "</colgroup>"
    pheads = "<th>Parameter</th><th>Value</th>" if has_params else ""
    notes = [n for n, _locked in card.notes if n.strip()]
    note_bars = "\n".join(f'<div class="note-bar">\U0001f4dd {_h(n)}</div>' for n in notes)
    note_section = "" if not notes else f"""<h2>Notes</h2>
{note_bars}
"""
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><title>{_h(card.id)} -- {_h(card.title)}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=B612:wght@400;700&display=swap" rel="stylesheet">
<link rel="stylesheet" href="http://{host}/static/starfire.css"></head>
<body class="print-doc">
<h1>{_h(card.title)}</h1>
<div class="meta">
  <span><b>ID:</b> {_h(card.id)}</span>
  <span><b>Rev:</b> {_h(card.revision)}</span>
  <span><b>Project:</b> {_h(card.project)}</span>
  <span><b>System/Subsystem:</b> {_h(card.system)} / {_h(card.subsystem)}</span>
  <span><b>Test Type:</b> {_h(card.test_type)}</span>
  <span><b>Requirement(s):</b> {_h(card.requirements)}</span>
  <span><b>Author:</b> {_h(card.author)}</span>
</div>
<div class="meta">
  <span><b>Test Conductor:</b> {_h(card.test_conductor)}</span>
  <span><b>Quality Assurance:</b> {_h(card.quality_assurance)}</span>
  <span><b>Date:</b> {_h(card.date)}</span>
</div>
{warn_section}{note_section}{nom_section}{ref_section}<h2>Test Points</h2>
<table class="test-points">
{cols}
<tr><th>#</th><th>Action</th>{pheads}<th>Expected Result</th><th>Observed Result</th><th>Reference</th>
<th>RP</th><th>Date</th><th>Time</th><th>Initial</th><th>Quality</th></tr>
{rows}
</table>
<script>window.onload = () => window.print();</script>
</body></html>"""


# Campaign data model
@dataclass
class CampaignTestRow:
    seq: int
    test_id: str = ""
    title: str = ""
    author: str = ""
    requirements: str = ""


@dataclass
class CampaignModule:
    # A script or binary in the program's dist/ folder that a test point's Activate button can run. 
    module: str = ""
    description: str = ""
    # "" = any test point; a card id (TC-001) = any point on that card;
    # card id + "-" + point number (TC-001-20) = that one point.
    scope: str = ""
    pinned: bool = False    # shown as an icon in the module column of the points in scope (first 4)


@dataclass
class Campaign:
    id: str = "SC-001"
    title: str = "Untitled Campaign"
    revision: str = "A"
    author: str = ""
    date: str = field(default_factory=lambda: date.today().isoformat())
    project: str = ""
    test_lead: str = ""
    quality_assurance: str = ""
    requirements: list = field(default_factory=list)   # [code, description, source, method]
    # Modules a test point's Activate button can run.
    modules: list = field(default_factory=list)         # list[CampaignModule]
    tests: list = field(default_factory=list)           # list[CampaignTestRow] -- "Test Series"
    ancillary_tests: list = field(default_factory=list)   # list[CampaignTestRow]
    sections: list = field(default_factory=list)         # list[(header, text)]
    jira_epic_key: str = ""
    reports: dict = field(default_factory=dict)
    # The campaign report (SR-xxx), once built: {"items": [section | vcrm | report | system, ...]}.
    campaign_report: dict = None
    live: list = field(default_factory=list)            # [{card, point, file}] Signet code can connect

    @staticmethod
    def from_dict(d: dict) -> "Campaign":
        row_known = {f.name for f in fields(CampaignTestRow)}   # drops the retired "status" key
        tests = [CampaignTestRow(**{k: v for k, v in t.items() if k in row_known}) for t in d.get("tests", [])]
        ancillary_tests = [CampaignTestRow(**{k: v for k, v in t.items() if k in row_known})
                           for t in d.get("ancillary_tests", [])]
        known = {f.name for f in fields(CampaignModule)}
        modules = [CampaignModule(**{k: v for k, v in m.items() if k in known}) for m in d.get("modules", [])]
        base = {k: v for k, v in d.items() if k not in ("tests", "ancillary_tests", "modules", "system", "subsystem")}
        return Campaign(tests=tests, ancillary_tests=ancillary_tests, modules=modules, **base)

    def to_dict(self) -> dict:
        return asdict(self)


STATUS_LABELS = {"TO_DO": "TO DO", "IN_PROGRESS": "IN PROGRESS", "COMPLETE": "COMPLETE"}

# A point's sign-off quality, worst to best. RFI (ready for issue) is above PASS.
RESULTS = ["FAIL", "CAUTION", "PASS", "RFI"]


def _quality(p) -> str:
    return (p.quality or "").upper()


def _is_run(p) -> bool:
    return bool((p.initial or "").strip()) and _quality(p) in RESULTS


def card_state(card: TestCard) -> dict:
    """Progress and worst result of a card, from its points' initials and quality.
    COMPLETE needs every point run and the last point RFI."""
    pts = sorted(card.points, key=lambda p: p.seq)
    run = [p for p in pts if _is_run(p)]
    if not run:
        progress = "TO_DO"
    elif len(run) == len(pts) and _quality(pts[-1]) == "RFI":
        progress = "COMPLETE"
    else:
        progress = "IN_PROGRESS"
    return {"progress": progress, "result": min((_quality(p) for p in run), key=RESULTS.index, default="")}


def status_label(state: dict) -> str:
    label = STATUS_LABELS[state["progress"]]
    return f"{label} ({state['result']})" if state["result"] else label


def card_warnings(card: TestCard) -> list:
    """Sign-off problems on one card, as short phrases."""
    pts = sorted(card.points, key=lambda p: p.seq)
    out = []
    for p in pts:
        initialed, q = bool((p.initial or "").strip()), _quality(p)
        if q in RESULTS and not initialed:
            out.append(f"point {p.seq} is {q} but has no initials")
        elif initialed and q not in RESULTS:
            out.append(f"point {p.seq} is initialed but has no PASS, CAUTION, FAIL or RFI")
        elif q in ("CAUTION", "FAIL"):
            out.append(f"point {p.seq} is {q}")
    if pts and all(_is_run(p) for p in pts) and _quality(pts[-1]) != "RFI":
        out.append("all points are run but the last is not RFI")
    elif pts and _is_run(pts[-1]) and _quality(pts[-1]) == "RFI" and not all(_is_run(p) for p in pts):
        out.append("the last point is RFI but earlier points are not run")
    return out


def _read_card(test_id: str):
    """The TestCard in the working directory for a Test ID, or None if missing or unreadable."""
    try:
        with open(os.path.join(os.getcwd(), test_filename(test_id)), "r", encoding="utf-8") as f:
            return TestCard.from_dict(json.loads(f.read()))
    except (OSError, json.JSONDecodeError, TypeError):
        return None


def campaign_statuses(camp: Campaign) -> dict:
    """Test ID -> computed status label for every row of both tests tables ("" if no .test file)."""
    out = {}
    for t in list(camp.tests) + list(camp.ancillary_tests):
        tid = (t.test_id or "").strip()
        if tid and tid not in out:
            card = _read_card(tid)
            out[tid] = status_label(card_state(card)) if card else ""
    return out


def next_test_id(camp: Campaign) -> str:
    """The next TC-nnn not used by a test row of the campaign or by a .test file here."""
    used = {(t.test_id or "").strip() for t in list(camp.tests) + list(camp.ancillary_tests)}
    n = 1
    while f"TC-{n:03d}" in used or os.path.exists(os.path.join(os.getcwd(), test_filename(f"TC-{n:03d}"))):
        n += 1
    return f"TC-{n:03d}"


def create_test(camp: Campaign, codes: list) -> dict:
    """Write an empty, untitled card that verifies the given requirements; returns its row fields."""
    known = {c for c, _, _, _ in requirement_rows(camp)}
    codes = [c for c in dict.fromkeys(str(c).strip() for c in codes) if c in known]
    if not codes:
        raise ValueError("no known requirement codes given")
    card = TestCard(id=next_test_id(camp), title="", author=camp.test_lead or camp.author, project=camp.project,
                    quality_assurance=camp.quality_assurance, requirements=", ".join(codes))
    with open(os.path.join(os.getcwd(), test_filename(card.id)), "w", encoding="utf-8") as f:
        json.dump(card.to_dict(), f, indent=2)
    return {"id": card.id, "title": "", "author": card.author, "requirements": card.requirements}


def card_text(text: str):
    """Normalized .test JSON if text is a card (a JSON object with a points list), else None."""
    try:
        d = json.loads(text)
        if not isinstance(d, dict) or not isinstance(d.get("points"), list):
            return None
        return json.dumps(TestCard.from_dict(d).to_dict(), indent=2)
    except (ValueError, TypeError, AttributeError, KeyError):
        return None


def _seq_int(seq) -> int:
    try:
        return int(float(str(seq).strip()))
    except ValueError:
        raise ValueError(f"point number {seq!r} is not a number")


def _new_points(extra: TestCard, seq) -> list:
    # A module card's points, renumbered by 10s from seq + 10; their attachments stay behind.
    start = _seq_int(seq)
    return [TestPoint(**{**asdict(p), "seq": start + 10 * (k + 1), "files": []}) for k, p in enumerate(extra.points)]


def _point_index(card: TestCard, seq) -> int:
    for i, p in enumerate(card.points):
        if str(p.seq).strip() == str(seq).strip():
            return i
    raise ValueError(f"no point {seq} on {card.id}")


def extend_card(card: TestCard, seq, extra: TestCard) -> TestCard:
    """The card with the module card's points added after point seq, which must be its last point."""
    if _point_index(card, seq) != len(card.points) - 1:
        raise ValueError("EXTEND works only on the last point")
    return TestCard.from_dict({**card.to_dict(), "points": [asdict(p) for p in card.points + _new_points(extra, seq)]})


def branch_card(card: TestCard, seq, extra: TestCard, new_id: str, campaign_id: str):
    """(new card, [(old, new) file renames]): this card up to and including point seq, under new_id,
    then the module card's points. Files named for the old card are renamed for the new one."""
    old, new = _point_prefix(campaign_id, card.id) + "-", _point_prefix(campaign_id, new_id) + "-"
    renames = []
    kept = []
    for p in card.points[:_point_index(card, seq) + 1]:
        files = []
        for n in p.files:
            if n.startswith(old):
                renames.append((n, new + n[len(old):]))
                n = new + n[len(old):]
            files.append(n)
        kept.append(TestPoint(**{**asdict(p), "files": files}))
    d = {**card.to_dict(), "id": new_id, "points": [asdict(p) for p in kept + _new_points(extra, seq)]}
    return TestCard.from_dict(d), renames


def feed_card(data: dict) -> dict:
    """BRANCH or EXTEND a card from the module-written .test file attached to one of its points.
    data: {"mode": "branch"|"extend", "card", "point", "file", "campaign", "used": [campaign test IDs]}."""
    card = TestCard.from_dict(data["card"])
    campaign_id = str(data.get("campaign") or "XX-000")
    name = data.get("file")
    if not (_own_file(campaign_id, card.id, name) and name.endswith(".test")):
        raise ValueError(f"{name!r} is not a .test file saved for this card")
    text = card_text(_read_text(_cwd_file(name), _POINT_FILE_MAX))
    if text is None:
        raise ValueError(f"{name} is not a test card")
    extra = TestCard.from_dict(json.loads(text))
    if data.get("mode") == "extend":
        return {"card": extend_card(card, data["point"], extra).to_dict()}
    if data.get("mode") != "branch":
        raise ValueError("mode must be branch or extend")
    branched, renames = branch_card(card, data["point"], extra, next_test_id(Campaign(tests=[CampaignTestRow(0, str(t)) for t in data.get("used") or []])), campaign_id)
    for a, b in renames:
        src, dst = _cwd_file(a), os.path.join(os.getcwd(), b)
        if src and not os.path.exists(dst):
            shutil.copy2(src, dst)
    with open(os.path.join(os.getcwd(), test_filename(branched.id)), "w", encoding="utf-8") as f:
        json.dump(branched.to_dict(), f, indent=2)
    return {"id": branched.id, "title": branched.title, "author": branched.author, "requirements": branched.requirements}


def test_filename(test_id: str) -> str:
    # The .test filename a campaign's Tests-table row links to, derived from Test ID rather than stored separately. 
    return _safe_filename(test_id, ".test")


def vcrm_filename(campaign_id: str) -> str:
    return _safe_filename(campaign_id, ".vcrm")


METHODS = ["T", "A", "D", "I"]   # Test, Analysis, Demonstration, Inspection


def requirement_rows(camp: Campaign) -> list:
    """(code, source, description, method) per requirement. Rows are stored as
    [code, description, source, method]; older files have only the first two."""
    out = []
    for r in camp.requirements:
        r = list(r) + [""] * (4 - len(r))
        out.append(((r[0] or "").strip(), r[2] or "", r[1] or "", r[3] or "T"))
    return out


def compute_vcrm(camp: Campaign) -> dict:
    """Verification Cross-Reference Matrix: requirement -> which Test
    Series tests verify it, with each test's computed status. One row per
    requirement and verifying test. Read-only and computed fresh each time
    from the campaign's Requirements table and each referenced test's own
    `requirements` field and point sign-offs. Ancillary Tests are excluded.
    """
    reqs = requirement_rows(camp)
    codes = [code for code, _, _, _ in reqs]
    cards: dict = {}        # test_id -> (card, state), or None if there is no readable .test file
    warnings = []
    for t in camp.tests:
        tid = (t.test_id or "").strip()
        if not tid or tid in cards:
            continue
        card = _read_card(tid)
        cards[tid] = (card, card_state(card)) if card else None
        if not card:
            warnings.append(f"{tid}: no {test_filename(tid)} file")
            continue
        tokens = [tok.strip() for tok in (card.requirements or "").split(",") if tok.strip()]
        if not tokens:
            warnings.append(f"{tid}: lists no requirements")
        warnings += [f"{tid}: {tok} is not in the Requirements table" for tok in tokens if tok not in codes]
        warnings += [f"{tid}: {g}" for g in card_warnings(card)]

    updated = datetime.now().strftime("%Y-%m-%d %H:%M")
    rows = []
    for req_code, source, req_desc, method in reqs:
        if not req_code:
            continue
        head = {"requirement": req_code, "source": source, "description": req_desc, "method": method}
        found = False
        for t in camp.tests:
            entry = cards.get((t.test_id or "").strip())
            if not entry:
                continue
            card, state = entry
            if req_code in [tok.strip() for tok in (card.requirements or "").split(",")]:
                found = True
                run = any(_is_run(p) for p in card.points)
                rows.append({**head, "event": f"{t.test_id} {card.title or t.title}".strip(),
                             "report": report_id(card.id) if run else "NONE", "status": status_label(state),
                             "updated": updated})
        if not found:
            rows.append({**head, "event": "NONE", "report": "NONE", "status": "--", "updated": updated})
            warnings.append(f"{req_code}: no test covers it")

    return {
        "campaign_id": camp.id, "campaign_title": camp.title,
        "generated": updated,
        "rows": rows, "warnings": warnings,
    }


# Export: print-ready HTML (same visual system as the test card's).
def export_campaign_print_html(camp: Campaign, host: str = "localhost") -> str:
    h = _h
    reqs = [r for r in requirement_rows(camp) if r[0] or r[2]]
    req_rows = "\n".join(f"<tr><td>{h(c)}</td><td>{h(src)}</td><td>{h(d)}</td><td>{h(m)}</td></tr>"
                         for c, src, d, m in reqs)
    req_section = "" if not reqs else f"""<h2>Requirements</h2>
<table><tr><th style="width:20%">Requirement</th><th style="width:16%">Source</th><th>Description</th><th style="width:8%">Method</th></tr>
{req_rows}
</table>
"""
    mods = [m for m in camp.modules if (m.module or "").strip()]
    mod_rows = "\n".join(
        f"<tr><td>{h(m.module)}</td><td>{h(m.description)}</td><td>{h(m.scope or 'any')}</td></tr>" for m in mods)
    mod_section = "" if not mods else f"""<h2>Modules</h2>
<table><tr><th style="width:25%">Module</th><th>Description</th><th style="width:15%">Scope</th></tr>
{mod_rows}
</table>
"""
    statuses = campaign_statuses(camp)
    test_rows = "\n".join(
        f"<tr><td>{t.seq}</td><td>{h(t.test_id)}</td><td>{h(t.title)}</td><td>{h(t.author)}</td>"
        f"<td>{h(t.requirements)}</td><td>{h(statuses.get((t.test_id or '').strip(), ''))}</td></tr>"
        for t in camp.tests
    )
    anc_rows = "\n".join(
        f"<tr><td>{t.seq}</td><td>{h(t.test_id)}</td><td>{h(t.title)}</td><td>{h(t.author)}</td>"
        f"<td>{h(t.requirements)}</td><td>{h(statuses.get((t.test_id or '').strip(), ''))}</td></tr>"
        for t in camp.ancillary_tests
    )
    anc_section = "" if not camp.ancillary_tests else f"""<h2>Ancillary Tests</h2>
<table>
<tr><th style="width:6%">#</th><th style="width:14%">Test ID</th><th>Title</th><th style="width:12%">Author</th><th style="width:16%">Requirements</th><th style="width:12%">Status</th></tr>
{anc_rows}
</table>
"""
    sections_html = "\n".join(
        f'<h2>{h(header)}</h2>\n<div class="section-text">{h(text)}</div>'
        for header, text in camp.sections
        if (header or "").strip() or (text or "").strip()
    )
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><title>{h(camp.id)} -- {h(camp.title)}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=B612:wght@400;700&display=swap" rel="stylesheet">
<link rel="stylesheet" href="http://{host}/static/starfire.css"></head>
<body class="print-doc">
<h1>{h(camp.title)}</h1>
<div class="meta">
  <span><b>ID:</b> {h(camp.id)}</span>
  <span><b>Rev:</b> {h(camp.revision)}</span>
  <span><b>Project:</b> {h(camp.project)}</span>
  <span><b>Author:</b> {h(camp.author)}</span>
</div>
<div class="meta">
  <span><b>Test Lead:</b> {h(camp.test_lead)}</span>
  <span><b>Quality Assurance:</b> {h(camp.quality_assurance)}</span>
  <span><b>Date:</b> {h(camp.date)}</span>
</div>
{req_section}{mod_section}<h2>Test Series</h2>
<table>
<tr><th style="width:6%">#</th><th style="width:14%">Test ID</th><th>Title</th><th style="width:12%">Author</th><th style="width:16%">Requirements</th><th style="width:12%">Status</th></tr>
{test_rows}
</table>
{anc_section}{sections_html}
<script>window.onload = () => window.print();</script>
</body></html>"""


VCRM_COLUMNS = [("Requirement", "requirement"), ("Source", "source"), ("Description", "description"),
                ("Method", "method"), ("Verification Event", "event"), ("Report", "report"), ("Status", "status"),
                ("Updated", "updated")]


def _vcrm_table_html(vcrm: dict) -> str:
    head = "".join(f"<th>{title}</th>" for title, _ in VCRM_COLUMNS)
    rows = "\n".join("<tr>" + "".join(f"<td>{_h(r.get(key, ''))}</td>" for _, key in VCRM_COLUMNS) + "</tr>"
                     for r in vcrm["rows"])
    return f"<table><tr>{head}</tr>\n{rows}\n</table>"


def export_vcrm_csv(vcrm: dict) -> str:
    # The VCRM as a Test Matrix: one row per requirement and verification event.
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([title for title, _ in VCRM_COLUMNS])
    for r in vcrm["rows"]:
        w.writerow([r.get(key, "") for _, key in VCRM_COLUMNS])
    return buf.getvalue()


def export_vcrm_print_html(vcrm: dict, host: str = "localhost") -> str:
    h = _h

    table = _vcrm_table_html(vcrm)
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><title>{h(vcrm['campaign_id'])} VCRM -- {h(vcrm['campaign_title'])}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=B612:wght@400;700&display=swap" rel="stylesheet">
<link rel="stylesheet" href="http://{host}/static/starfire.css"></head>
<body class="print-doc">
<h1>Verification Cross-Reference Matrix</h1>
{table}
<script>window.onload = () => window.print();</script>
</body></html>"""


# Export: .docx (python-docx).
_HEADER_FILL = "E7EEF5"
_WARN_FILL = "F8D7D7"
_NOTE_FILL = "E7EEF5"


def _shade_cell(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _add_table(doc: Document, headers: list, rows: list, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, htext in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = htext
        cell.paragraphs[0].runs[0].bold = True
        _shade_cell(cell, _HEADER_FILL)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return table


def _add_callout(doc: Document, text: str, hex_fill: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.rows[0].cells[0]
    cell.text = text
    _shade_cell(cell, hex_fill)
    # No visible grid for callouts -- clear the borders python-docx's
    # default table style would otherwise draw.
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def export_campaign_docx(camp: Campaign) -> bytes:
    doc = Document()
    doc.add_heading(camp.title, level=0)
    meta = doc.add_paragraph()
    meta.add_run(
        f"ID: {camp.id}    Rev: {camp.revision}    Project: {camp.project}    Author: {camp.author}\n"
        f"Test Lead: {camp.test_lead}    Quality Assurance: {camp.quality_assurance}    "
        f"Date: {camp.date}"
    ).font.size = Pt(9.5)

    reqs = [r for r in requirement_rows(camp) if r[0] or r[2]]
    if reqs:
        doc.add_heading("Requirements", level=1)
        _add_table(doc, ["Requirement", "Source", "Description", "Method"], reqs)

    mods = [m for m in camp.modules if (m.module or "").strip()]
    if mods:
        doc.add_heading("Modules", level=1)
        _add_table(doc, ["Module", "Description", "Scope"],
                   [(m.module, m.description, m.scope or "any") for m in mods])

    statuses = campaign_statuses(camp)
    doc.add_heading("Test Series", level=1)
    _add_table(
        doc, ["#", "Test ID", "Title", "Author", "Requirements", "Status"],
        [(t.seq, t.test_id, t.title, t.author, t.requirements, statuses.get((t.test_id or "").strip(), ""))
         for t in camp.tests],
    )

    if camp.ancillary_tests:
        doc.add_heading("Ancillary Tests", level=1)
        _add_table(
            doc, ["#", "Test ID", "Title", "Author", "Requirements", "Status"],
            [(t.seq, t.test_id, t.title, t.author, t.requirements, statuses.get((t.test_id or "").strip(), ""))
             for t in camp.ancillary_tests],
        )

    for header, text in camp.sections:
        if not (header or "").strip() and not (text or "").strip():
            continue
        doc.add_heading(header, level=1)
        for line in (text or "").split("\n"):
            doc.add_paragraph(line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _vcrm_docx_table(doc, vcrm: dict) -> None:
    _add_table(doc, [title for title, _ in VCRM_COLUMNS],
               [[r.get(key, "") for _, key in VCRM_COLUMNS] for r in vcrm["rows"]])


def export_vcrm_docx(vcrm: dict) -> bytes:
    doc = Document()
    doc.add_heading("Verification Cross-Reference Matrix", level=0)
    _vcrm_docx_table(doc, vcrm)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_test_docx(card: "TestCard") -> bytes:
    doc = Document()
    doc.add_heading(card.title, level=0)
    meta = doc.add_paragraph()
    meta.add_run(
        f"ID: {card.id}    Rev: {card.revision}    Project: {card.project}    "
        f"System/Subsystem: {card.system} / {card.subsystem}    Test Type: {card.test_type}    "
        f"Author: {card.author}\n"
        f"Test Conductor: {card.test_conductor}    Quality Assurance: {card.quality_assurance}    "
        f"Date: {card.date}    Requirement(s): {card.requirements}"
    ).font.size = Pt(9.5)

    warnings = [w for w, _locked in card.warnings if w.strip()]
    if warnings:
        doc.add_heading("Warnings", level=1)
        for w in warnings:
            _add_callout(doc, "\u26a0 " + w, _WARN_FILL)

    notes = [n for n, _locked in card.notes if n.strip()]
    if notes:
        doc.add_heading("Notes", level=1)
        for n in notes:
            _add_callout(doc, "\U0001f4dd " + n, _NOTE_FILL)

    nom = sorted_nomenclature(card)
    if nom:
        doc.add_heading("Nomenclature", level=1)
        _add_table(
            doc, ["Symbol", "Definition", "Units"],
            [(symbol_display(s), d, u) for s, d, u, _locked in nom],
        )

    refs = [(r, d) for r, d, _locked in card.references if (r or "").strip() or (d or "").strip()]
    if refs:
        doc.add_heading("References", level=1)
        _add_table(doc, ["Reference", "Description"], refs)

    doc.add_heading("Test Points", level=1)
    has_params = any(_point_params(p)[0] for p in card.points)
    _add_table(
        doc,
        ["#", "Action"] + (["Parameter", "Value"] if has_params else []) +
        ["Expected Result", "Observed Result", "Reference", "Responsible Party", "Date", "Time", "Initial", "Quality"],
        [(p.seq, p.action, *(_point_params(p) if has_params else ()), p.expected, p.observed, p.reference,
          p.responsible_party, p.date, p.time, p.initial, p.quality) for p in card.points],
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# Test Card (procedure only) vs Test Report (as-run record).
def procedure_only(card: TestCard) -> TestCard:
    """A copy with every run result cleared, so the card prints as a blank procedure."""
    d = card.to_dict()
    d["article"] = ""
    for p in d["points"]:
        for k in ("observed", "responsible_party", "date", "time", "initial", "quality"):
            p[k] = ""
        p["images"], p["files"] = [], []
    return TestCard.from_dict(d)


def report_id(card_id: str) -> str:
    cid = (card_id or "").strip()
    return "TR-" + cid[3:] if cid.upper().startswith("TC-") else "TR-" + (cid or "000")



def report_head(card: TestCard) -> str:
    rev = (card.revision or "").strip()
    return " ".join(x for x in (report_id(card.id), (card.title or "").strip(), f"REV {rev.upper()}" if rev else "") if x)


def sr_id(campaign_id: str) -> str:
    cid = (campaign_id or "").strip()
    return "SR-" + cid[3:] if cid.upper().startswith("SC-") else "SR-" + (cid or "000")


def sr_head(camp: Campaign) -> str:
    rev = (camp.revision or "").strip()
    return " ".join(x for x in (sr_id(camp.id), (camp.title or "").strip(), f"REV {rev.upper()}" if rev else "") if x)


def _sr_items(camp: Campaign):
    # Yield (kind, payload) for each item that has something to show; a test with no saved .test file is skipped.
    items = (camp.campaign_report or {}).get("items", [])
    for it in sorted(items, key=lambda x: {"report": 1, "system": 2}.get(x.get("type"), 0)):   # then Test Reports, then Systems
        kind = it.get("type")
        if kind == "section":
            idx = it.get("idx")
            if idx is not None:   # Live reference to the campaign's Sections
                secs = camp.sections or []
                it = {"header": secs[idx][0], "text": secs[idx][1]} if isinstance(idx, int) and 0 <= idx < len(secs) else {}
            if (it.get("header") or "").strip() or (it.get("text") or "").strip():
                yield kind, it
        elif kind == "vcrm":
            yield kind, compute_vcrm(camp)
        elif kind in ("report", "system"):
            path = os.path.join(os.getcwd(), test_filename(str(it.get("card") or "")))
            try:
                with open(path, encoding="utf-8") as f:
                    yield kind, TestCard.from_dict(json.load(f))
            except (OSError, ValueError):
                continue


def export_campaign_report_html(camp: Campaign, host: str = "localhost", preview: bool = False) -> str:
    head = sr_head(camp)
    parts = []
    for kind, v in _sr_items(camp):
        if kind == "section":
            parts.append(f'<div class="sr-part"><h2>{_h(v.get("header"))}</h2><div class="section-text">{_h(v.get("text"))}</div></div>')
        elif kind == "vcrm":
            parts.append(f'<div class="sr-part sr-new"><h2>Verification Cross-Reference Matrix</h2>{_vcrm_table_html(v)}</div>')
        elif kind == "system":
            parts.append(f'<div class="sr-part sr-new">{system_html(system_graph(v.to_dict(), camp.id, camp.live))}</div>')
        else:
            parts.append(f'<div class="sr-part sr-new">{report_body(v, camp.id, (camp.reports or {}).get(v.id, {}), camp.live)}</div>')
    body = f"""<h1>{_h(head)}</h1>
<div class="meta">
  <span><b>Campaign:</b> {_h(camp.id)}</span>
  <span><b>Project:</b> {_h(camp.project)}</span>
</div>
<div class="meta">
  <span><b>Test Lead:</b> {_h(camp.test_lead)}</span>
  <span><b>Quality Assurance:</b> {_h(camp.quality_assurance)}</span>
  <span><b>Date:</b> {_h(camp.date)}</span>
</div>
{"".join(parts)}"""
    return _report_page(head, body, host, preview)


def _add_field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def export_campaign_report_docx(camp: Campaign) -> bytes:
    doc = Document()
    foot = doc.sections[0].footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run("Page ")
    _add_field(foot, "PAGE")
    foot.add_run(" of ")
    _add_field(foot, "NUMPAGES")
    doc.add_heading(sr_head(camp), level=0)
    doc.add_paragraph().add_run(
        f"Campaign: {camp.id}    Project: {camp.project}\n"
        f"Test Lead: {camp.test_lead}    Quality Assurance: {camp.quality_assurance}    Date: {camp.date}"
    ).font.size = Pt(9.5)
    for kind, v in _sr_items(camp):
        if kind == "section":
            doc.add_heading(v.get("header") or "", level=1)
            for line in (v.get("text") or "").split("\n"):
                doc.add_paragraph(line)
            continue
        doc.add_page_break()
        if kind == "vcrm":
            doc.add_heading("Verification Cross-Reference Matrix", level=1)
            _vcrm_docx_table(doc, v)
        elif kind == "system":
            doc.add_heading(f"System {sy_id(v.id)}", level=1)
            doc.add_paragraph("The System diagram is in the PDF export.")
        else:
            _report_docx_body(doc, v, camp.id, (camp.reports or {}).get(v.id, {}))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif"}


def _file_kind(name: str) -> str:
    ext = os.path.splitext(name)[1].lower()
    if ext == ".txt":
        return "csv" if _looks_csv(_cwd_file(name)) else "txt"
    return "csv" if ext == ".csv" else "image" if ext in _IMAGE_TYPES else "other"


def _point_prefix(campaign_id: str, card_id: str) -> str:
    part = lambda s, d: re.sub(r"[^A-Za-z0-9._-]", "_", str(s or "").strip()) or d
    return part(campaign_id, "XX-000") + "-" + part(card_id, "TC")


def _own_file(campaign_id: str, card_id: str, name) -> bool:
    # A bare name of an existing file starting <campaign>-<card>-.
    return (isinstance(name, str) and name == os.path.basename(name)
            and name.startswith(_point_prefix(campaign_id, card_id) + "-") and os.path.isfile(os.path.join(os.getcwd(), name)))


def _point_file_pattern(campaign_id: str, card_id: str):
    return re.compile(r"^" + re.escape(_point_prefix(campaign_id, card_id)) + r"-(\d+)(?:-[^/]*)?\.(?!tmp\Z)[A-Za-z0-9]+$")   # not a .tmp: a helper's file before its rename


def list_point_files(campaign_id: str, card_id: str, attached: dict = None) -> dict:
    """{seq: [{"name", "kind"}]}. With `attached` ({seq: [names]}, a point's own list) it
    returns those that are files of this card. Without it, every file in the working folder
    named <campaign>-<card>-<point>[-anything].<ext>, to attach files that predate the list."""
    out = {}
    if attached is not None:
        for seq, names in attached.items():
            for n in names if isinstance(names, list) else []:
                if _own_file(campaign_id, card_id, n):
                    f = {"name": n, "kind": _file_kind(n)}
                    if f["kind"] == "csv":
                        f["cols"] = csv_y_columns(_cwd_file(n))
                        f["ydef"] = csv_y_default(_cwd_file(n))
                    out.setdefault(str(seq), []).append(f)
        return out
    pat = _point_file_pattern(campaign_id, card_id)
    try:
        names = sorted(os.listdir(os.getcwd()))
    except OSError:
        return out
    for n in names:
        m = pat.match(n)
        if m and os.path.isfile(n):
            out.setdefault(m.group(1), []).append({"name": n, "kind": _file_kind(n)})
    return out


_CARD_FILE_MAX = 25_000_000   # bytes


def add_card_file(campaign_id: str, card_id: str, point, name: str, data: bytes) -> dict:
    """Save an uploaded file as <campaign>-<card>-<point>-<name> (a name already starting
    <campaign>-<card>- is kept); never replaces an existing file with different contents."""
    prefix = _point_prefix(campaign_id, card_id) + "-"
    seq = str(point)
    base = re.sub(r"[^A-Za-z0-9._-]", "_", os.path.basename(str(name or "")).strip()).lstrip(".")
    target = base if base.startswith(prefix) else f"{prefix}{seq}-{base}"
    if not seq.isdigit() or not base or len(target) > 200:
        return {"saved": False, "error": "bad file name"}
    if len(data) > _CARD_FILE_MAX:
        return {"saved": False, "error": f"file is over {_CARD_FILE_MAX // 1_000_000} MB"}
    path = os.path.join(os.getcwd(), target)
    try:
        with open(path, "xb") as f:
            f.write(data)
    except FileExistsError:
        # Picking a file already in the folder (same name, same bytes) just attaches it.
        with open(path, "rb") as f:
            if f.read(_CARD_FILE_MAX + 1) == data:
                return {"saved": True, "filename": target}
        return {"saved": False, "error": f"{target} already exists with different contents"}
    except OSError as e:
        return {"saved": False, "error": str(e)}
    return {"saved": True, "filename": target}


def _num(s):
    m = re.match(r"\s*([-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?)", str(s))
    return float(m.group(1)) if m else None


def _csv_table(path: str, limit: int = 5000):
    """(header, body rows, numeric column indices) of a csv; blank rows and # lines are skipped."""
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        rows = [r for r in csv.reader(f) if any(c.strip() for c in r) and not r[0].lstrip().startswith("#")][:limit + 1]
    if not rows:
        return None
    header = rows[0] if any(_num(c) is None for c in rows[0]) else []
    body = rows[1:] if header else rows
    width = max(len(r) for r in body) if body else 0
    return header, body, [i for i in range(width) if any(i < len(r) and _num(r[i]) is not None for r in body)]


def _default_y(header: list, numeric: list) -> int:
    # The "value" column if there is one, else the numeric column after x.
    names = [c.strip().lower() for c in header]
    return names.index("value") if "value" in names and names.index("value") in numeric else numeric[min(1, len(numeric) - 1)]


def _col_label(header: list, i: int, default: str) -> str:
    return header[i] if header and i < len(header) and header[i].strip() else default


def csv_series(path: str, limit: int = 5000, y: int = None):
    """(x_label, y_label, xs, ys) from a csv: y is the given column, else the "value" column if there is
    one, else the second numeric column; x is the first numeric column, else the row number."""
    t = _csv_table(path, limit)
    if not t or not t[2]:
        return None
    header, body, numeric = t
    if y not in numeric:
        y = _default_y(header, numeric)
    x = numeric[0] if numeric[0] != y else None
    pts = []
    for i, r in enumerate(body):
        yv = _num(r[y]) if y < len(r) else None
        xv = (_num(r[x]) if x < len(r) else None) if x is not None else i + 1
        if yv is not None and xv is not None:
            pts.append((xv, yv))
    if not pts:
        return None
    return (_col_label(header, x, "x") if x is not None else "sample", _col_label(header, y, "y"),
            [p[0] for p in pts], [p[1] for p in pts])


def csv_y_default(path: str):
    # The column a plot traces when none is chosen.
    t = _csv_table(path, 50) if path else None
    return _default_y(t[0], t[2]) if t and t[2] else None


def csv_y_columns(path: str) -> list:
    # [[index, label]] of the columns a plot can trace: every numeric column after the first (x).
    t = _csv_table(path, 50) if path else None
    return [[i, _col_label(t[0], i, f"column {i + 1}")] for i in t[2][1:]] if t else []


def _looks_csv(path) -> bool:
    # A .txt that is comma-separated with at least two plottable rows ("# ..." lines are ignored).
    try:
        s = csv_series(path, 50) if path else None
        first = next(l for l in _read_text(path, 4096).splitlines() if l.strip() and not l.lstrip().startswith("#"))
    except (OSError, StopIteration):
        return False
    return bool(s) and len(s[2]) >= 2 and "," in first


def csv_plot_svg(path: str, cols: list = None) -> str:
    """One trace (csv_series' default column), or up to four of csv_y_columns' (cols), against the
    first numeric column; several traces get classes s0-s3 (colored in starfire.css) and a legend."""
    ok = {i for i, _ in csv_y_columns(path)}
    cols = sorted({c for c in (cols or []) if c in ok})[:4] or [None]
    return _plot_svg([s for s in (csv_series(path, y=c) for c in cols) if s])


def _plot_svg(series: list) -> str:
    # Up to four (x_label, y_label, xs, ys) traces on one set of axes.
    if not series:
        return ""
    multi = len(series) > 1
    allx = [v for s in series for v in s[2]]
    ally = [v for s in series for v in s[3]]
    x0, x1 = min(allx), max(allx)
    lo, hi = min(ally), max(ally)
    pad = (hi - lo) * 0.1 or 1.0
    y0, y1 = lo - pad, hi + pad
    sx = lambda v: 56 + (v - x0) / ((x1 - x0) or 1) * 484
    sy = lambda v: 176 - (v - y0) / (y1 - y0) * 156
    g = lambda v: f"{v:.4g}"
    traces = ""
    for k, (_, _, xs, ys) in enumerate(series):
        cls = f' class="s{k}"' if multi else ""
        line = " ".join(f"{sx(a):.1f},{sy(b):.1f}" for a, b in zip(xs, ys))
        dots = "".join(f'<circle{cls} cx="{sx(a):.1f}" cy="{sy(b):.1f}" r="3"/>' for a, b in zip(xs, ys)) if len(xs) <= 60 else ""
        traces += f'<polyline{cls} points="{line}"/><g>{dots}</g>'
    xl, yl = series[0][0], "" if multi else series[0][1]
    legend = ('<text x="540" y="12" text-anchor="end">' + "".join(
        f'<tspan class="s{k}"{' dx="12"' if k else ""}>{_h(s[1])}</tspan>' for k, s in enumerate(series)) + "</text>") if multi else ""
    label = ", ".join(s[1] for s in series)
    return (f'<svg class="rp-plot" viewBox="0 0 560 210" role="img" aria-label="{_h(label)} vs {_h(xl)}">'
            f'<line x1="56" y1="176" x2="540" y2="176"/><line x1="56" y1="20" x2="56" y2="176"/>'
            f'<text x="50" y="24" text-anchor="end">{g(y1)}</text><text x="50" y="176" text-anchor="end">{g(y0)}</text>'
            f'<text x="56" y="194">{g(x0)}</text><text x="540" y="194" text-anchor="end">{g(x1)}</text>'
            f'<text x="298" y="206" text-anchor="middle">{_h(xl)}</text>'
            f'<text x="12" y="98" transform="rotate(-90 12 98)" text-anchor="middle">{_h(yl)}</text>'
            f'{legend}{traces}</svg>')


def _cwd_file(name: str):
    path = os.path.join(os.getcwd(), os.path.basename(name or ""))
    return path if name and os.path.isfile(path) else None


def _img_tag(name: str) -> str:
    path = _cwd_file(name)
    mime = _IMAGE_TYPES.get(os.path.splitext(name)[1].lower())
    if not path or not mime:
        return f'<div class="rp-missing">{_h(name)} (not found)</div>'
    with open(path, "rb") as f:
        data = base64.b64encode(f.read()).decode()
    return f'<figure class="rp-img"><img src="data:{mime};base64,{data}" alt="{_h(name)}"><figcaption>{_h(name)}</figcaption></figure>'


def _read_text(path: str, limit: int = 200_000) -> str:
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read(limit)


def _ticked(files: dict, card: TestCard, campaign_id: str, seq) -> list:
    # Ticked names for a point, kept only if the point still has them attached.
    own = next((p.files for p in card.points if str(p.seq) == str(seq)), [])
    return [n for n in (files or {}).get(str(seq), []) if n in own and _own_file(campaign_id, card.id, n)]


def _plotted(files: dict, seq) -> list:
    # Names the user chose to plot; stored under files["plot"] as {seq: [names]}.
    p = (files or {}).get("plot")
    return p.get(str(seq), []) if isinstance(p, dict) else []


def _plot_cols(files: dict, seq, name: str) -> list:
    # Y columns chosen for a plotted csv; stored under files["cols"] as {seq: {name: [indices]}}.
    c = ((files or {}).get("cols") or {}).get(str(seq)) or {}
    v = c.get(name) if isinstance(c, dict) else None
    return [i for i in v if isinstance(i, int)] if isinstance(v, list) else None


def _validate_fails(path) -> "int | None":
    # validate.py output (first line "# validate"): its FAIL line count; None for any other file.
    try:
        lines = [l for l in _read_text(path, None).splitlines() if l.strip()] if path else []
    except OSError:
        return None
    return sum(l.startswith("FAIL ") for l in lines) if lines and lines[0].strip() == "# validate" else None


def _validate_warning(name: str) -> str:
    n = _validate_fails(_cwd_file(name))
    return f"LIMIT EXCEEDED: {n} limit{'s' if n > 1 else ''} failed ({name})" if n else ""


def _txt_pre(path) -> str:
    # Text output; a validate.py file's FAIL lines are wrapped for the amber style.
    text = _cap_lines(_read_text(path))
    if _validate_fails(path) is None:
        return f'<pre class="rp-out">{_h(text)}</pre>'
    return '<pre class="rp-out">' + "\n".join(
        f'<span class="rp-fail">{_h(l)}</span>' if l.startswith("FAIL ") else _h(l) for l in text.split("\n")) + "</pre>"


def _point_file_html(name: str, plot: bool = False, cols: list = None) -> str:
    kind, path = _file_kind(name), _cwd_file(name)
    if kind == "csv" and not plot:
        return f'<div class="rp-file"><div class="rp-cap">{_h(name)}</div><pre class="rp-out">{_h(_cap_lines(_read_text(path)))}</pre></div>'
    if kind == "csv":
        svg = csv_plot_svg(path, cols)
        return f'<div class="rp-file"><div class="rp-cap">{_h(name)}</div>{svg or "<div class=rp-missing>no numeric data to plot</div>"}</div>'
    if kind == "image":
        return _img_tag(name)
    if kind == "txt":
        return f'<div class="rp-file"><div class="rp-cap">{_h(name)}</div>{_txt_pre(path)}</div>'
    return f'<div class="rp-file rp-cap">Attached: {_h(name)}</div>'


def report_body(card: TestCard, campaign_id: str = "", files: dict = None, live=None) -> str:
    # One Test Report as an HTML fragment (header through test points, then its System if added).
    pts = []
    for p in card.points:
        signed = _h(p.initial)
        when = _h(f"{p.date} {p.time}".strip())
        under = "".join(f'<span class="rp-when">{x}</span>' for x in (when, signed) if x)
        imgs = "".join(_img_tag(n) for n in p.images)
        ticked = _ticked(files, card, campaign_id, p.seq)
        extra = "".join(_point_file_html(n, n in _plotted(files, p.seq), _plot_cols(files, p.seq, n)) for n in ticked)
        warn = "".join(f'<div class="rp-warn">{_h(w)}</div>' for w in map(_validate_warning, ticked) if w)
        q = (p.quality or "").upper()
        pts.append(f'''<section class="rp-point rp-b-{_h(q.lower()) or "none"}" data-seq="{p.seq}">
<div class="rp-head"><span>{p.seq} &nbsp; {_h(p.action)}</span><span class="rp-right"><span class="rp-q rp-q-{_h(q.lower())}">{_h(q) or "&mdash;"}</span>{under}</span></div>
<div class="rp-line"><b>Expected:</b> {_h(p.expected)}</div>
<div class="rp-line"><b>Observed:</b> {_h(p.observed) or "&mdash;"}</div>{f'<div class="rp-line rp-params"><b>Parameters:</b> {_h(_params_line(p))}</div>' if _params_line(p) else ""}{warn}
{imgs and f'<div class="rp-imgs">{imgs}</div>'}{extra}
</section>''')
    head = report_head(card)
    nom = sorted_nomenclature(card)
    nom_html = "" if not nom else "<h2>Nomenclature</h2><table><tr><th>Symbol</th><th>Definition</th><th>Units</th></tr>" + "".join(
        f"<tr><td>{_h(symbol_display(a))}</td><td>{_h(b)}</td><td>{_h(u)}</td></tr>" for a, b, u, _l in nom) + "</table>"
    refs = [(a, b) for a, b, _l in card.references if (a or "").strip() or (b or "").strip()]
    ref_html = "" if not refs else "<h2>References</h2><table><tr><th>Reference</th><th>Description</th></tr>" + "".join(
        f"<tr><td>{_h(a)}</td><td>{_h(b)}</td></tr>" for a, b in refs) + "</table>"
    return f"""<h1>{_h(head)}</h1>
<div class="meta">
  <span><b>Campaign:</b> {_h(campaign_id or "XX-000")}</span>
  <span><b>Article:</b> {_h(card.article) or "&mdash;"}</span>
  <span><b>System/Subsystem:</b> {_h(card.system)} / {_h(card.subsystem)}</span>
</div>
<div class="meta">
  <span><b>Test Conductor:</b> {_h(card.test_conductor)}</span>
  <span><b>Quality Assurance:</b> {_h(card.quality_assurance)}</span>
  <span><b>Date:</b> {_h(card.date)}</span>
</div>
<h2>Requirements covered</h2>
<p>{_h(card.requirements) or "&mdash;"}</p>
{nom_html}{ref_html}<h2>Test points</h2>
{"".join(pts)}{system_html(system_graph(card.to_dict(), campaign_id, live or [])) if (files or {}).get("system") else ""}"""


def _report_page(title: str, body: str, host: str, preview: bool) -> str:
    script = "" if preview else "<script>window.onload = () => window.print();</script>"
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><title>{_h(title)}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=B612:wght@400;700&family=B612+Mono:wght@400;700&display=swap" rel="stylesheet">
<link rel="stylesheet" href="http://{host}/static/starfire.css"></head>
<body class="print-doc report-doc{' preview' if preview else ''}">
{body}
{script}
</body></html>"""


def export_report_html(card: TestCard, campaign_id: str = "", files: dict = None,
                       host: str = "localhost", preview: bool = False, live=None) -> str:
    return _report_page(report_head(card), report_body(card, campaign_id, files, live), host, preview)


def _report_docx_body(doc, card: TestCard, campaign_id: str = "", files: dict = None) -> None:
    doc.add_heading(report_head(card), level=0)
    doc.add_paragraph().add_run(
        f"Campaign: {campaign_id or 'XX-000'}    "
        f"Article: {card.article}    System/Subsystem: {card.system} / {card.subsystem}\nTest Conductor: {card.test_conductor}    "
        f"Quality Assurance: {card.quality_assurance}    Date: {card.date}"
    ).font.size = Pt(9.5)
    doc.add_heading("Requirements covered", level=1)
    doc.add_paragraph(card.requirements or "--")
    nom = sorted_nomenclature(card)
    if nom:
        doc.add_heading("Nomenclature", level=1)
        _add_table(doc, ["Symbol", "Definition", "Units"], [(symbol_display(a), b, u) for a, b, u, _l in nom])
    refs = [(a, b) for a, b, _l in card.references if (a or "").strip() or (b or "").strip()]
    if refs:
        doc.add_heading("References", level=1)
        _add_table(doc, ["Reference", "Description"], refs)
    doc.add_heading("Test points", level=1)
    for p in card.points:
        doc.add_heading(f"{p.seq}  {p.action}    [{(p.quality or '--').upper()}]", level=2)
        doc.add_paragraph(f"Expected: {p.expected}\nObserved: {p.observed or '--'}")
        params = _params_line(p)
        if params:   # one line, cut like the PDF's two-line limit
            doc.add_paragraph("Parameters: " + (params if len(params) <= 200 else params[:200] + "..."))
        for n in list(p.images) + _ticked(files, card, campaign_id, p.seq):
            path, kind = _cwd_file(n), _file_kind(n)
            if kind == "image" and path:
                doc.add_picture(path, width=Inches(3.5))
                doc.add_paragraph(n).runs[0].font.size = Pt(8)
            elif kind == "csv" and path and n not in _plotted(files, p.seq):
                doc.add_paragraph(n).runs[0].font.size = Pt(8)
                run = doc.add_paragraph().add_run(_cap_lines(_read_text(path), 200))
                run.font.name, run.font.size = "Courier New", Pt(8)
            elif kind == "csv" and path:
                s = csv_series(path)
                doc.add_paragraph(f"{n} (plot in the PDF export)").runs[0].font.size = Pt(8)
                if s:
                    _add_table(doc, [s[0], s[1]], [(f"{a:g}", f"{b:g}") for a, b in list(zip(s[2], s[3]))[:50]])
            elif kind == "txt" and path:
                if _validate_warning(n):
                    w = doc.add_paragraph().add_run(_validate_warning(n))
                    w.bold, w.font.color.rgb = True, RGBColor(0xB3, 0x7D, 0x00)
                doc.add_paragraph(n).runs[0].font.size = Pt(8)
                run = doc.add_paragraph().add_run(_cap_lines(_read_text(path), 200))
                run.font.name, run.font.size = "Courier New", Pt(8)
            else:
                doc.add_paragraph(f"Attached: {n}" if path else f"{n} (not found)")
        sign = [x for x in (f"{p.date} {p.time}".strip(), p.initial) if x]
        if sign:
            doc.add_paragraph("    ".join(sign)).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if (files or {}).get("system"):
        doc.add_heading(f"System {sy_id(card.id)}", level=1)
        doc.add_paragraph("The System diagram is in the PDF export.")


def export_report_docx(card: TestCard, campaign_id: str = "", files: dict = None) -> bytes:
    doc = Document()
    _report_docx_body(doc, card, campaign_id, files)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()



# Browser UI






# HTTP server

def _safe_filename(stem: str, ext: str) -> str:
    """Turn a card id into a filesystem-safe filename, confined to the
    current directory (no path separators from user-controlled input)."""
    stem = os.path.basename((stem or "card").strip()) or "card"
    stem = re.sub(r"[^A-Za-z0-9._-]", "_", stem)
    return stem + ext




# Optional integrations, each checked once at startup: jira_config.json turns Jira tickets on.


def _program_dir() -> str:
    return os.path.dirname(os.path.abspath(sys.executable if getattr(sys, "frozen", False) else __file__))


def _modules_dir() -> str:
    # The dist/ folder: next to the built binary, or ./dist beside this script when run from source (where starfi.sh puts its output).
    here = _program_dir()
    return here if getattr(sys, "frozen", False) else os.path.join(here, "dist")


# Jira tickets: on only if jira_config.json is in the working directory or next to the program.
JIRA_CONFIG_FILENAME = "jira_config.json"

JIRA_CONFIG_TEMPLATE = {
    "base_url": "https://yourteam.atlassian.net",
    "email": "you@example.com",
    "api_token": "",
    "project_key": "PROJ",
    # Bug for test point tickets, Task for campaign test rows.
    "issue_type": "Bug",
    "test_issue_type": "Task",
    # Set automatically after the first ticket with an epic key.
    "epic_link_field": None,
}


def _jira_config_path() -> str:
    # Working directory first, then beside the program (dist/ when run from source).
    for d in (os.getcwd(), _program_dir(), _modules_dir()):
        p = os.path.join(d, JIRA_CONFIG_FILENAME)
        if os.path.isfile(p):
            return p
    return os.path.join(os.getcwd(), JIRA_CONFIG_FILENAME)


def jira_load_config() -> dict | None:
    path = _jira_config_path()
    if not os.path.isfile(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as f:
            cfg = json.load(f)
    except (OSError, json.JSONDecodeError):
        return None
    required = ("base_url", "email", "api_token", "project_key")
    if not all(cfg.get(k) for k in required):
        return None
    cfg["issue_type"] = cfg.get("issue_type") or "Bug"
    cfg["test_issue_type"] = cfg.get("test_issue_type") or "Task"
    cfg["epic_link_field"] = cfg.get("epic_link_field") or None
    return cfg


def _jira_save_epic_field(field_name: str) -> None:
    # Best-effort cache; a failed write just means re-detecting next time.
    path = _jira_config_path()
    try:
        with open(path, "r", encoding="utf-8") as f:
            cfg = json.load(f)
        cfg["epic_link_field"] = field_name
        with open(path, "w", encoding="utf-8") as f:
            json.dump(cfg, f, indent=2)
    except (OSError, json.JSONDecodeError):
        pass


def _jira_auth_header(cfg: dict) -> str:
    token = base64.b64encode(f"{cfg['email']}:{cfg['api_token']}".encode()).decode()
    return f"Basic {token}"


def _jira_api_request(cfg: dict, method: str, path: str, body: dict | None = None) -> tuple[int, dict]:
    url = cfg["base_url"].rstrip("/") + path
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, method=method)
    req.add_header("Authorization", _jira_auth_header(cfg))
    req.add_header("Content-Type", "application/json")
    req.add_header("Accept", "application/json")
    try:
        with urllib.request.urlopen(req, timeout=10) as resp:
            raw = resp.read()
            return resp.status, (json.loads(raw) if raw else {})
    except urllib.error.HTTPError as e:
        raw = e.read()
        try:
            return e.code, json.loads(raw) if raw else {}
        except json.JSONDecodeError:
            return e.code, {"errorMessages": [raw.decode(errors="replace")]}
    except urllib.error.URLError as e:
        return 0, {"errorMessages": [str(e.reason)]}


def _jira_error_text(body: dict) -> str:
    parts = list(body.get("errorMessages") or [])
    for field, msg in (body.get("errors") or {}).items():
        parts.append(f"{field}: {msg}")
    return "; ".join(parts) if parts else "unknown Jira error"


def _jira_find_epic_field(cfg: dict) -> str | None:
    """Classic projects: find the custom field named 'Epic Link'."""
    status, body = _jira_api_request(
        cfg, "GET",
        f"/rest/api/2/issue/createmeta?projectKeys={cfg['project_key']}&expand=projects.issuetypes.fields",
    )
    if status != 200:
        return None
    for project in body.get("projects", []):
        for issuetype in project.get("issuetypes", []):
            for field_id, field_meta in issuetype.get("fields", {}).items():
                if (field_meta.get("name") or "").strip().lower() == "epic link":
                    return field_id
    return None


def jira_create_ticket(summary: str, description: str, epic_key: str = "",
                    issue_type: str | None = None) -> dict:
    """Returns {"success": True, "key", "url"} or {"success": False, "error"}; never raises.
    issue_type overrides the config's issue_type for this call."""
    cfg = jira_load_config()
    if not cfg:
        return {"success": False,
                 "error": f"No usable {JIRA_CONFIG_FILENAME} found "
                          f"(or it's missing base_url/email/api_token/project_key)."}

    issue = {
        "project": {"key": cfg["project_key"]},
        "summary": summary,
        "description": description,
        "issuetype": {"name": issue_type or cfg["issue_type"]},
    }

    epic_key = (epic_key or "").strip()
    if epic_key:
        cached = cfg.get("epic_link_field")
        if cached:
            issue[cached] = epic_key if cached != "parent" else {"key": epic_key}
        else:
            # Team-managed projects first; no metadata lookup needed.
            issue["parent"] = {"key": epic_key}

    status, body = _jira_api_request(cfg, "POST", "/rest/api/2/issue", {"fields": issue})

    if status == 201:
        if epic_key and not cfg.get("epic_link_field"):
            _jira_save_epic_field("parent")
        key = body.get("key", "")
        return {"success": True, "key": key, "url": f"{cfg['base_url'].rstrip('/')}/browse/{key}"}

    # "parent" failed: likely a classic project, so retry with the Epic Link field.
    if epic_key and not cfg.get("epic_link_field") and "parent" in issue:
        custom_field = _jira_find_epic_field(cfg)
        if custom_field:
            issue.pop("parent")
            issue[custom_field] = epic_key
            status2, body2 = _jira_api_request(cfg, "POST", "/rest/api/2/issue", {"fields": issue})
            if status2 == 201:
                _jira_save_epic_field(custom_field)
                key = body2.get("key", "")
                return {"success": True, "key": key,
                         "url": f"{cfg['base_url'].rstrip('/')}/browse/{key}"}
            return {"success": False, "error": _jira_error_text(body2)}

    return {"success": False, "error": _jira_error_text(body)}


JIRA_AVAILABLE = jira_load_config() is not None


def _static_dir() -> str:
    # Beside this script; in a PyInstaller build the bundled data is in sys._MEIPASS instead.
    base = sys._MEIPASS if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS") else os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, "static")


def _read_static(name: str) -> str:
    with open(os.path.join(_static_dir(), name), encoding="utf-8") as f:
        return f.read()


# Modules: scripts/binaries the user drops into dist/ and registers in a
# campaign's Modules table; a test point's Activate button runs one.

# Starfire's own files live in dist/ too; none of them is a module.
_RESERVED_MODULES = {"starfire"}
# A file with one of these extensions is run by its interpreter, so it needs
# no execute bit or shebang. Anything else must be an executable file (a
# compiled binary, or a script with a shebang and chmod +x) and is run directly.
_INTERPRETERS = {".py": "python3", ".sh": "bash", ".jl": "julia", ".signet": ""}   # .signet runs inside Starfire, no interpreter
_MODULE_TIMEOUT_S = 3600      # a module may legitimately run a whole test
_MODULE_OUTPUT_LINES = 4096   # output lines kept; a longer run keeps its first and last half
_RUNS: dict = {}              # (card, point) -> running Popen, so BREAK can find it
_BROKEN: set = set()          # runs a BREAK was sent to
_RUNS_LOCK = threading.Lock()


def _cap_lines(text: str, limit: int = None) -> str:
    limit = limit or _MODULE_OUTPUT_LINES
    lines = text.splitlines(keepends=True)
    if len(lines) <= limit:
        return text
    half = limit // 2
    omitted = len(lines) - 2 * half
    return "".join(lines[:half]) + f"... {omitted} lines omitted ...\n" + "".join(lines[-half:])


def _is_module(d: str, name: str) -> bool:
    path = os.path.join(d, name)
    if name in _RESERVED_MODULES or name.startswith(".") or not os.path.isfile(path):
        return False
    ext = os.path.splitext(name)[1]
    return ext in _INTERPRETERS or (signet and ext.lower() in signet.SOURCE_TYPES) or os.access(path, os.X_OK)


def list_modules() -> list:
    """Names of runnable files in dist/: regular, not hidden, not Starfire's
    own, and either an interpreted type (.py, .sh, .jl, .signet), a source Signet builds (.cpp, .rs, .adb, .f90)
    or executable."""
    d = _modules_dir()
    try:
        names = os.listdir(d)
    except OSError:
        return []
    return sorted(n for n in names if _is_module(d, n))


def _module_env(card_id, point_seq, campaign_id="", params=(), stem="", inputs_path="") -> dict:
    env = dict(os.environ, STARFIRE_CAMPAIGN=str(campaign_id or "XX-000"),
               STARFIRE_CARD=str(card_id), STARFIRE_POINT=str(point_seq))
    # SIGNET_ and STARFIRE_ PREFIX (names files for this point), MODULE and SIGNALS_IN, and the helpers
    prefix = f"{_point_prefix(campaign_id, card_id)}-{point_seq}"
    if signet:
        env = signet.module_env(env, prefix, stem, inputs_path)
    else:
        env.update(STARFIRE_PREFIX=prefix, STARFIRE_MODULE=stem)
    for name, value in params:
        env["STARFIRE_PARAM_" + name.upper()] = value
    if getattr(sys, "frozen", False):
        # The one-file bundle points LD_LIBRARY_PATH at its own libraries;
        # a system python3 or a compiled module inheriting that can crash on
        # a library mismatch. PyInstaller keeps the original in *_ORIG.
        if "LD_LIBRARY_PATH_ORIG" in env:
            env["LD_LIBRARY_PATH"] = env["LD_LIBRARY_PATH_ORIG"]
        else:
            env.pop("LD_LIBRARY_PATH", None)
    return env


_MODULE_ARGS_MAX = 4096


def _parse_args(text) -> list:
    """Split a command-line string the way a POSIX shell would (quotes, backslashes),
    without running one. Newlines are whitespace; lines starting with # are skipped.
    Raises ValueError on bad quoting, NUL or an oversized line."""
    text = "" if text is None else str(text)
    if len(text) > _MODULE_ARGS_MAX or "\0" in text:
        raise ValueError("arguments too long or contain a NUL byte")
    text = text.replace("\\\n", "")  # line continuation, as in a shell
    text = "\n".join(ln for ln in text.splitlines() if not ln.lstrip().startswith("#"))
    return shlex.split(text)


_PARAM_NAME = re.compile(r"[A-Za-z_][A-Za-z0-9_]*\Z")


def _parse_params(rows) -> list:
    """A point's parameters as [(name, value)]. Rows with a blank name are skipped.
    Raises ValueError on a bad or repeated name (ignoring case), NUL or an oversized value."""
    out, seen = [], set()
    for row in rows or []:
        if not isinstance(row, (list, tuple)) or len(row) < 2:
            raise ValueError("expected [name, value] rows")
        name, value = str(row[0] or "").strip(), "" if row[1] is None else str(row[1])
        if not name:
            continue
        if not _PARAM_NAME.match(name):
            raise ValueError(f"{name!r}: a name is letters, digits and _, not starting with a digit")
        if name.lower() in seen:
            raise ValueError(f"{name!r} is listed twice")
        if len(value) > _MODULE_ARGS_MAX or "\0" in value:
            raise ValueError(f"{name!r}: value too long or contains a NUL byte")
        seen.add(name.lower())
        out.append((name, value))
    return out


# Signals on points: the host Signet code runs against (see README.md, Signals and Signet; the language is in signet.py)
def _live_index(campaign_id: str, live) -> list:
    """The live list as [(card, point, file, display)], keeping files that exist and are named for their point."""
    out = []
    for e in live or []:
        if not isinstance(e, dict):
            continue
        card, point, name = str(e.get("card", "")), str(e.get("point", "")), e.get("file")
        if not isinstance(name, str) or name != os.path.basename(name) or not os.path.isfile(name):
            continue
        m = _point_file_pattern(campaign_id, card).match(name)
        if not m or m.group(1) != point:
            continue
        head = f"{_point_prefix(campaign_id, card)}-{point}-"
        out.append((card, point, name, name[len(head):] if name.startswith(head) else name))
    return out


def _live_match(text: str, card_id: str, index: list):
    """The live entry (card, point, file, display) for [point:][module.]port.sig or [point:]file, or an error message."""
    card = point = None
    rest = text
    if ":" in text:
        head, rest = text.split(":", 1)
        m = re.match(r"(?:(.+)-)?(\d+)\Z", head)
        if not m:
            return f"{head!r} isn't a point"
        card, point = m.group(1) or card_id, m.group(2)
    if rest.endswith(".sig"):
        want = rest[:-4].split(".")
        hit = lambda d: d.endswith(".sig") and len(d[:-4].split(".")) == 2 and d[:-4].split(".")[-len(want):] == want
    else:
        hit = lambda d: d == rest
    found = [e for e in index if hit(e[3]) and (card is None or e[0].lower() == card.lower())
             and (point is None or e[1] == point)]
    if not found:
        return f"nothing live matches {text}"
    if len(found) > 1:
        return f"{text} matches {len(found)} live signals: " + ", ".join(f"{c}-{p}:{d}" for c, p, _, d in found)
    return found[0]


class _PointHost:
    """What Signet code on a point sees: the campaign's live list and the point's parameters."""

    def __init__(self, card_id, campaign_id, live, params, pts=None):
        self.card, self.index = str(card_id), _live_index(campaign_id, live)
        self.params = {k.lower(): v for k, v in params}
        self.pts = pts if isinstance(pts, dict) else {}   # the card's points, for a sequence box
        self.lines = []    # echo output, shown above the module's own
        self.point = None  # set for a module's own run, so input() can find the point's sequence

    def ask(self, prompt, type_name):
        # input(): only a sequence has someone to ask (its point's page); the point's own run shares its answers
        with _RUNS_LOCK:
            r = (_JOBS.get((self.card, self.point)) or {}).get("runner")   # the job's, so a sequence that ended still answers
        if r is None:
            raise LookupError("input needs a sequence: in Starfire, only a sequence asks, on its point's page")
        return r.ask("run", prompt, type_name)

    def resolve(self, text):
        r = _live_match(text, self.card, self.index)
        if isinstance(r, str):
            raise LookupError(r)
        c, p, name, disp = r
        return signet.source_value(name, f"{c}-{p}:{disp}", text)

    def param(self, name):
        if name.lower() not in self.params:
            raise LookupError(f"this point has no parameter {name}")
        return self.params[name.lower()]

    def point_param(self, point, name):
        info = self.pts.get(str(point))
        if not isinstance(info, dict):
            raise LookupError(f"no point {point} on this card")
        for row in info.get("params") or []:
            if isinstance(row, (list, tuple)) and len(row) > 1 and str(row[0]).strip().lower() == name.lower():
                return "" if row[1] is None else str(row[1])
        raise LookupError(f"point {point} has no parameter {name}")

    def autorun(self):
        return False    # off in Starfire 4: run(n) always waits for RUN on the point

    def echo(self, text):
        self.lines.append(text)


# Systems: a card's Signet code drawn as boxes and lines (see README.md, Systems)
_SY_STATE = {"failed": 2, "ok": 1, "": 0}    # box outline: red, blue, dim blue


def sy_id(card_id: str) -> str:
    cid = (card_id or "").strip()
    return "SY-" + cid[3:] if cid.upper().startswith("TC-") else "SY-" + (cid or "000")


def _static_resolve(text: str, card_id: str, index: list):
    """A ref's live entry (card, point, file, display) without reading it, or an error string."""
    try:
        return _live_match(text, card_id, index)
    except re.error:
        return f"can't read {text}"


# Port map: each point's declared ports and the point-to-point connections in the Signals boxes,
# for the connection checkboxes on the test page (see README.md, Connecting points)
_POINT_SIG = re.compile(r"^(?:(\d+):)?([A-Za-z0-9_-]+)\.([A-Za-z_][A-Za-z0-9_]*)\.sig$")


def _connectable(src, dst) -> bool:
    """Whether an output declared `src` can be an input declared `dst`'s one source: signet's check, and a
    vector or series only to its own kind (signet lets a single value into a vector input)."""
    if src and dst and ({src[0], dst[0]} & {"vector", "series"}) and src[0] != dst[0]:
        return False
    return signet._fits(src, dst)


def _mergeable(src, dst) -> bool:
    """Whether an output declared `src` can be one of several merged into an input declared `dst`."""
    if not dst:
        return True
    if dst[0] == "vector":
        return signet._fits(src, dst[1:]) or signet._fits(src, dst)
    return dst[0] == "series" and src == dst


def _ref_point(text, rows, i):
    """(seq, output) of an earlier point on the card that a ref like 10:psu.volts.sig names, or None."""
    m = _POINT_SIG.match(text)
    if not m:
        return None
    seq, stem, port = m.groups()
    hits = [r for r in rows[:i] if r["stem"] == stem and (seq is None or r["seq"] == seq)]
    return (hits[0]["seq"], port) if len(hits) == 1 else None


def port_map(points) -> dict:
    """points: the card's points in order, [{seq, module, signals}]. {"points": [{seq, module, stem, ports, error,
    sequence}], "links": [{src, out, dst, input, box, line, bad}], "pairs": [{src, out, dst, input, fits, merge}]}: links
    are the connections from earlier points the Signals boxes make now (box: whose code), pairs the ones the
    declared ports allow."""
    if not signet:
        return {"points": [], "links": [], "pairs": [], "error": _NO_SIGNET}
    pts = [p for p in points or [] if isinstance(p, dict)]
    rows = []
    for p in pts:
        name = str(p.get("module") or "").strip()
        row = {"seq": str(p.get("seq")), "module": name, "stem": os.path.splitext(name)[0], "ports": None, "error": "",
               "sequence": False}
        if name:
            try:
                row["ports"] = _ports_lookup([name])(row["stem"])
            except ValueError as e:
                row["error"] = str(e)
        rows.append(row)
    links, where = [], {r["seq"]: k for k, r in enumerate(rows)}
    for i, (r, p) in enumerate(zip(rows, pts)):
        code = str(p.get("signals") or "")
        if not code.strip():
            continue
        try:
            r["sequence"] = signet.is_sequence(code)
            conns = signet.connections(code, points=True)
            merged = {nd[3][1]: [nd[3][1]] + [e[1] for e in nd[4] if e[0] == "ref"]
                      for nd in signet._walk(signet.parse_signals(code)) if nd[0] == "route" and nd[2] == "merge"}
        except ValueError:
            continue    # the run reports it
        for c in conns:
            if c["module"]:            # this point's own module, connect(x, multiply.volts)
                j = i if c["module"] == r["stem"] else None
            else:                      # another point's, connect(x, 20:volts): a sequence's go through its run()
                j = None if r["sequence"] else where.get(c["point"])
            if j is None:
                continue
            for ref in c["refs"]:
                for text in merged.get(ref, [ref]):
                    src = _ref_point(text, rows, j)
                    link = src and {"src": src[0], "out": src[1], "dst": rows[j]["seq"], "input": c["port"],
                                    "box": r["seq"], "line": c["line"]}
                    if link and link not in links:
                        links.append(link)
    for l in links:       # a connection the declared ports don't allow (written in Signet code, not by the page)
        a, b = rows[where[l["src"]]], rows[where[l["dst"]]]
        outs, ins = (a["ports"] or {}).get("out"), (b["ports"] or {}).get("in")
        l["bad"] = ""
        if a["ports"] and l["out"] not in (outs or {}):
            l["bad"] = f"{a['module']} has no output {l['out']}"
        elif b["ports"] and l["input"] not in (ins or {}):
            l["bad"] = f"{b['module']} has no input {l['input']}"
        elif outs and ins:
            ow, nw = outs[l["out"]], ins[l["input"]]
            n = sum(1 for k in links if k["dst"] == l["dst"] and k["input"] == l["input"])
            if not (signet._fits(ow, nw) or _mergeable(ow, nw)):
                l["bad"] = f"{l['src']}:{l['out']} is {'::'.join(ow)}, but {l['dst']}:{l['input']} takes {'::'.join(nw)}"
            elif n > 1 and (not nw or nw[0] not in ("vector", "series")):
                l["bad"] = f"{l['dst']}:{l['input']} takes one source, not {n}"
    pairs = []
    for i, a in enumerate(rows):
        for o, ow in ((a["ports"] or {}).get("out") or {}).items():
            for b in rows[i + 1:]:
                if not signet.STEM.match(b["stem"]):
                    continue
                for n, nw in ((b["ports"] or {}).get("in") or {}).items():
                    f, m = _connectable(ow, nw), _mergeable(ow, nw)
                    if f or m:
                        pairs.append({"src": a["seq"], "out": o, "dst": b["seq"], "input": n, "fits": f, "merge": m})
    return {"points": rows, "links": links, "pairs": pairs}


def system_graph(card: dict, campaign_id: str, live) -> dict:
    """A card's System as it is now: boxes (modules on points, sources on other cards, live files,
    multiplex buses) and one line per source of every connect() in its Signet code. A box is
    outlined by its module's last run; a line is amber when its source has data. No signal values."""
    card_id = str(card.get("id") or "")
    if not signet:
        return {"format": "starfire-system", "version": 1, "system": sy_id(card_id), "card": card_id,
                "campaign": campaign_id or "XX-000", "generated": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
                "boxes": [], "lines": [], "errors": [_NO_SIGNET]}
    points = {str(p.get("seq")): p for p in card.get("points", []) if isinstance(p, dict)}
    index = _live_index(campaign_id, live)
    boxes, lines, errors = {}, [], []
    other_cards = {}

    def box(key, **kw):
        if key not in boxes:
            boxes[key] = dict(id=f"b{len(boxes)}", inputs=[], outputs=[], **kw)
        return boxes[key]

    def port(lst, name, **kw):
        for p in lst:
            if p["port"] == name:
                return p
        lst.append(dict(port=name, **kw))
        return lst[-1]

    def run_state(c, pt, stem):
        # The module's last run on that point: "failed" (or broken), "ok", or "" (not run).
        if c.lower() == card_id.lower():
            p = points.get(pt) or {}
        else:
            if c not in other_cards:
                oc = _read_card(c)
                other_cards[c] = {str(q.seq): {"runs": q.runs} for q in oc.points} if oc else {}
            p = other_cards[c].get(pt) or {}
        run = (p.get("runs") or {}).get(stem)
        return "failed" if run in ("failed", "broken") else "ok" if run == "complete" else ""

    def mark(b, state):
        if _SY_STATE[state] > _SY_STATE[b.get("state", "")]:
            b["state"] = state

    # Modules with current .sig outputs on this card's points
    head = _point_prefix(campaign_id, card_id) + "-"
    out_pat = re.compile(r"^" + re.escape(head) + r"(\d+)-([A-Za-z0-9_-]+)\.([A-Za-z_][A-Za-z0-9_]*)\.sig$")
    try:
        names_on_disk = sorted(os.listdir(os.getcwd()))
    except OSError:
        names_on_disk = []
    for n in names_on_disk:
        m = out_pat.match(n)
        if m and m.group(1) in points:
            b = box(("m", m.group(1), m.group(2)), kind="module", point=m.group(1), module=m.group(2),
                    label=f"{m.group(1)}:{m.group(2)}")
            port(b["outputs"], m.group(3), file=n)

    for seq in sorted(points, key=lambda s: int(s) if s.isdigit() else 0):
        code = points[seq].get("signals") or ""
        if not code.strip():
            continue
        try:
            conns = signet.connections(code, points=not signet.is_sequence(code))   # a sequence's go through run()
        except ValueError as e:
            errors += [f"point {seq}: {ln}" for ln in str(e).splitlines()]
            continue
        def yield_connect(cn):
            line_no, words, stem, inp = cn["line"], cn["types"], cn["module"], cn["port"]
            dseq = seq
            if not stem:    # aimed at another point: connect(x, 20:volts)
                dseq = cn["point"]
                stem = os.path.splitext(str((points.get(dseq) or {}).get("module") or "").strip())[0]
                if not stem:
                    return
            dst = box(("m", dseq, stem), kind="module", point=dseq, module=stem, label=f"{dseq}:{stem}")
            refs, consts, buses = cn["refs"], cn["constants"], cn["buses"]
            ip = port(dst["inputs"], inp)
            if consts and not refs and not buses:
                ip["const"] = ", ".join(consts)
            def source_box(ref):
                r = _static_resolve(ref, card_id, index)
                if isinstance(r, str):
                    errors.append(f"point {seq} line {line_no}: {r}")
                    b = box(("x", ref), kind="missing", label=ref, dim=True)
                    port(b["outputs"], ref)
                    return b, ref, "dim", None, ref
                c, pt, fname, disp = r
                src = f"{c}-{pt}:{disp}"
                if disp.endswith(".sig"):
                    stem2, prt = disp[:-4].split(".")
                    if c.lower() == card_id.lower() and pt in points:
                        b = box(("m", pt, stem2), kind="module", point=pt, module=stem2, label=f"{pt}:{stem2}")
                        port(b["outputs"], prt, file=fname)
                    else:
                        b = box(("e", c.upper(), pt), kind="external", label=f"{c}-{pt}", dim=True)
                        mark(b, run_state(c, pt, stem2))
                        prt = disp
                        port(b["outputs"], prt, file=fname)
                    try:
                        signet.read_sig(fname)
                        color = "amber"
                    except ValueError:
                        color = "dim"
                    return b, prt, color, fname, src
                b = box(("f", c.upper(), pt, disp), kind="file", label=f"{pt}:{disp}" if c.lower() == card_id.lower() else src,
                        dim=c.lower() != card_id.lower())
                port(b["outputs"], disp, file=fname)
                return b, disp, "amber", fname, src
            def add_line(frm, fport, color, fname, src, to, tport, bus=False):
                lines.append({"from": frm["id"], "from_port": fport, "to": to["id"], "to_port": tport,
                              "color": color, "bus": bus, "file": fname, "source": src,
                              "type": "::".join(words), "point": seq, "line": line_no})
            for ref in refs:
                b, prt, color, fname, src = source_box(ref)
                add_line(b, prt, color, fname, src, dst, inp)
            for bundle in buses:
                jb = box(("j", seq, line_no, inp), kind="bus", label="mux")
                worst = "amber"
                for fld, rs in bundle:
                    port(jb["inputs"], fld)
                    for ref in rs:
                        b, prt, color, fname, src = source_box(ref)
                        add_line(b, prt, color, fname, src, jb, fld)
                        worst = "dim" if color == "dim" else worst
                port(jb["outputs"], "bus")
                add_line(jb, "bus", worst, None, "multiplex", dst, inp, bus=True)
        for cn in conns:
            yield_connect(cn)
    for b in boxes.values():
        if b["kind"] == "module":
            mark(b, run_state(card_id, b["point"], b["module"]))

    return {"format": "starfire-system", "version": 1, "system": sy_id(card_id), "card": card_id,
            "campaign": campaign_id or "XX-000",
            "generated": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
            "boxes": list(boxes.values()), "lines": lines, "errors": errors}


def signal_value(name: str) -> dict:
    """A live .sig's value for the System view: its type, the value as text if small, and a plot for a vector or series."""
    name = os.path.basename(str(name or ""))
    if not signet:
        return {"error": _NO_SIGNET}
    if not name.endswith(".sig") or not os.path.isfile(name):
        return {"error": "not found"}
    try:
        obj = signet.read_sig(name)
    except ValueError as e:
        return {"error": f"{name} {e}"}
    v, t = obj["value"], str(obj.get("type") or signet.infer_type(obj["value"]))
    out = {"type": t, "origin": obj.get("origin", ""), "generated": obj.get("generated", "")}
    series = []
    s = signet.series_parts(v) if isinstance(v, dict) else None
    if s:
        vals = v["values"]
        x = v["x"] if "x" in v else [v["start"] + i * v["step"] for i in range(len(vals["re"] if isinstance(vals, dict) else vals))]
        xl = {"time": "time (s)", "frequency": "frequency (Hz)"}.get(s[0], s[0])
        if isinstance(vals, dict):
            series = [(xl, "re", x, vals["re"]), (xl, "im", x, vals["im"])]
        else:
            series = [(xl, "value", x, vals)]
    elif isinstance(v, list) and v and all(isinstance(r, list) for r in v):
        series = [("sample", f"row {k + 1}", list(range(1, len(r) + 1)), r) for k, r in enumerate(v)]
    elif isinstance(v, list) and len(v) > 1:
        series = [("sample", "value", list(range(1, len(v) + 1)), v)]
    ok = lambda x, y: signet.is_num(x) and signet.is_num(y)
    series = [(a, b, [x for x, y in zip(xs, ys) if ok(x, y)], [y for x, y in zip(xs, ys) if ok(x, y)]) for a, b, xs, ys in series]
    series = [q for q in series if q[2]]
    if series:
        out["svg"] = _plot_svg(series[:4])
    h = _head(v)
    if h:
        out["head"], out["count"] = h
    text = json.dumps(v)
    out["text"] = text if len(text) <= 400 else text[:400] + " ..."
    out["scalar"] = not isinstance(v, (list, dict))
    return out

def _head(v, n: int = 5):
    # The first few entries of a vector or series as text, and how many there are.
    f = lambda q: f"{q:.6g}" if isinstance(q, float) else json.dumps(q)
    s = signet.series_parts(v) if isinstance(v, dict) else None
    if s:
        vals = v["values"]
        ys = [complex(a, b) for a, b in zip(vals["re"], vals["im"])] if isinstance(vals, dict) else vals
        xs = v["x"] if "x" in v else [v["start"] + i * v["step"] for i in range(n)]
        g = lambda q: f"{q.real:.6g}{q.imag:+.6g}i" if isinstance(q, complex) else f(q)
        return [f"{f(x)}: {g(y)}" for x, y in zip(xs[:n], ys[:n])], len(ys)
    if isinstance(v, list):
        return [f(q) for q in v[:n]], len(v)
    return None


def point_signals(card: dict, campaign_id: str, seq, live) -> dict:
    """The .sig files a point's Measure column can show: its modules' outputs, then the
    signals its Signet code connects in. Unsaved code counts, as in a System."""
    if not signet:
        return {"signals": [], "error": _NO_SIGNET}
    card_id, seq = str(card.get("id") or ""), str(seq)
    out, seen = [], set()
    pat = re.compile(r"^" + re.escape(f"{_point_prefix(campaign_id, card_id)}-{seq}-")
                     + r"([A-Za-z0-9_-]+)\.([A-Za-z_][A-Za-z0-9_]*)\.sig$")
    try:
        names = sorted(os.listdir(os.getcwd()))
    except OSError:
        names = []
    for n in names:
        m = pat.match(n)
        if m:
            out.append({"file": n, "label": f"{m.group(1)}.{m.group(2)}", "role": "out"})
            seen.add(n)
    p = next((q for q in card.get("points", []) if isinstance(q, dict) and str(q.get("seq")) == seq), {})
    code = p.get("signals") or ""
    if not code.strip():
        return {"signals": out}
    try:
        conns = signet.connections(code)
    except ValueError as e:
        return {"signals": out, "error": str(e).splitlines()[0]}
    index = _live_index(campaign_id, live)
    for cn in conns:
        refs = list(cn["refs"]) + [ref for bundle in cn["buses"] for _, rs in bundle for ref in rs]
        for ref in refs:
            r = _static_resolve(ref, card_id, index)
            if isinstance(r, str) or not r[2].endswith(".sig") or r[2] in seen:
                continue
            seen.add(r[2])
            out.append({"file": r[2], "label": f'{cn["module"]}.{cn["port"]} ← {r[0]}-{r[1]}:{r[3]}', "role": "in"})
    return {"signals": out}


_SY_BW, _SY_BUSW, _SY_HEAD, _SY_ROW, _SY_GAPX, _SY_GAPY, _SY_PAD = 200, 70, 26, 18, 110, 28, 24


def system_svg(g: dict) -> str:
    """The System diagram as SVG. Columns follow the data: each box one column right of the
    furthest box feeding it; a line that skips columns runs through a slot under each one."""
    boxes = [dict(b) for b in g.get("boxes", [])]
    by_id = {b["id"]: b for b in boxes}
    lines = g.get("lines", [])
    layer = {b["id"]: 0 for b in boxes}
    for _ in range(len(boxes)):
        moved = False
        for ln in lines:
            if ln["from"] != ln["to"] and layer[ln["to"]] < layer[ln["from"]] + 1:
                layer[ln["to"]] = layer[ln["from"]] + 1
                moved = True
        if not moved:
            break
    cols = {}
    for b in boxes:
        b["w"] = _SY_BUSW if b["kind"] == "bus" else _SY_BW
        b["h"] = _SY_HEAD + max(len(b["inputs"]), len(b["outputs"]), 1) * _SY_ROW + 8
        cols.setdefault(layer[b["id"]], []).append(b)
    vias = []
    for ln in lines:
        v = []
        for c in range(layer[ln["from"]] + 1, layer[ln["to"]]):
            slot = {"slot": True, "w": _SY_BW, "h": 10, "label": "", "kind": "slot"}
            cols.setdefault(c, []).append(slot)
            v.append(slot)
        vias.append(v)
    def key(b):
        p = str(b.get("point") or "")
        return (2 if b.get("slot") else 0 if b["kind"] == "module" else 1, int(p) if p.isdigit() else 0, b["label"])
    width = height = 0
    for c, col in cols.items():
        col.sort(key=key)
        y = _SY_PAD
        for b in col:
            b["x"], b["y"] = _SY_PAD + c * (_SY_BW + _SY_GAPX), y
            y += b["h"] + _SY_GAPY
        width = max(width, _SY_PAD + c * (_SY_BW + _SY_GAPX) + _SY_BW + _SY_PAD)
        height = max(height, y + _SY_PAD)
    width, height = max(width, 300), max(height, 120)

    def port_y(b, lst, name):
        i = next((k for k, p in enumerate(lst) if p["port"] == name), 0)
        return b["y"] + _SY_HEAD + i * _SY_ROW + _SY_ROW / 2

    out = []
    for i, (ln, via) in enumerate(zip(lines, vias)):
        a, b = by_id[ln["from"]], by_id[ln["to"]]
        pts = [(a["x"] + a["w"], port_y(a, a["outputs"], ln["from_port"]))]
        for v in via:
            pts += [(v["x"], v["y"] + v["h"] / 2), (v["x"] + v["w"], v["y"] + v["h"] / 2)]
        pts.append((b["x"], port_y(b, b["inputs"], ln["to_port"])))
        d = f"M{pts[0][0]:.0f},{pts[0][1]:.0f}"
        for k in range(1, len(pts)):
            (ax, ay), (bx, by) = pts[k - 1], pts[k]
            if k % 2 == 0 and via:      # straight across a slot
                d += f" L{bx:.0f},{by:.0f}"
                continue
            dx = max(40, (bx - ax) / 2) if bx > ax else 90
            d += f" C{ax + dx:.0f},{ay:.0f} {bx - dx:.0f},{by:.0f} {bx:.0f},{by:.0f}"
        cls = "sy-line " + ln["color"] + (" bus" if ln.get("bus") else "")
        tip = f'{ln["source"]} → {b["label"]}.{ln["to_port"]}' + (f' ({ln["type"]})' if ln.get("type") else "")
        out.append(f'<g class="sy-link" data-line="{i}"><path class="sy-hit" d="{d}"/><path class="{cls}" d="{d}"/>'
                   f'<title>{_h(tip)}</title></g>')
    for b in by_id.values():
        x, y, w = b["x"], b["y"], b["w"]
        parts = [f'<g class="sy-box {b["kind"]}{" dim" if b.get("dim") else ""}{" " + b["state"] if b.get("state") else ""}">'
                 f'<rect x="{x}" y="{y}" width="{w}" height="{b["h"]}" rx="5"/>'
                 f'<text class="sy-label" x="{x + 8}" y="{y + 17}">{_h(b["label"])}</text>']
        for k, p in enumerate(b["inputs"]):
            py = y + _SY_HEAD + k * _SY_ROW + _SY_ROW / 2
            const = f' <tspan class="sy-const">= {_h(p["const"])}</tspan>' if p.get("const") else ""
            parts.append(f'<circle class="sy-port" cx="{x}" cy="{py:.0f}" r="3"/>'
                         f'<text class="sy-port-name" x="{x + 8}" y="{py + 4:.0f}">{_h(p["port"])}{const}</text>')
        if b["kind"] != "bus":
            for k, p in enumerate(b["outputs"]):
                py = y + _SY_HEAD + k * _SY_ROW + _SY_ROW / 2
                parts.append(f'<circle class="sy-port" cx="{x + w}" cy="{py:.0f}" r="3"/>'
                             f'<text class="sy-port-name" x="{x + w - 8}" y="{py + 4:.0f}" text-anchor="end">{_h(p["port"])}</text>')
        out.append("".join(parts) + "</g>")
    if not out:
        out.append(f'<text class="sy-port-name" x="{_SY_PAD}" y="{_SY_PAD + 10}">No connections or signals on this card yet.</text>')
    return (f'<svg class="sy-diagram" xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" '
            f'width="{width}" height="{height}" role="img" aria-label="{_h(g.get("system", "System"))}">{"".join(out)}</svg>')


def system_signals_html(g: dict) -> str:
    """The expanded value of each vector or series a System connects: its type, plot and text.
    A scalar's value is already in the connection table."""
    out, seen = [], set()
    for ln in g.get("lines", []):
        f = str(ln.get("file") or "")
        if not f.endswith(".sig") or f in seen:
            continue
        seen.add(f)
        v = signal_value(f)
        if "error" in v or v.get("scalar"):
            continue
        out.append(f'<div class="sy-sig"><div class="sy-sig-head">{_h(ln["source"])} '
                   f'<span>{_h(v["type"])} | {_h(v.get("generated", ""))}</span></div>'
                   f'{v.get("svg", "")}<pre>{_h(v["text"])}</pre></div>')
    return f'<h3>Signals</h3>{"".join(out)}' if out else ""


def system_html(g: dict) -> str:
    """A System for a report: heading, diagram, the connection table (each source's type, as
    connected else as stored, and value: a scalar's, else [...]) and the expanded signals."""
    by_id = {b["id"]: b for b in g.get("boxes", [])}
    rows = []
    for ln in g.get("lines", []):
        typ, val = ln.get("type", ""), ""
        if ln.get("file", "") and str(ln["file"]).endswith(".sig"):
            v = signal_value(ln["file"])
            if "error" not in v:
                typ, val = typ or v.get("type", ""), v.get("text", "") if v.get("scalar") else "[...]"
        rows.append((ln["source"], f'{by_id[ln["to"]]["label"]}.{ln["to_port"]}', typ, val, ln["color"]))
    for b in g.get("boxes", []):
        rows += [(p["const"], f'{b["label"]}.{p["port"]}', "", "", "") for p in b["inputs"] if p.get("const")]
    cells = []
    for src, inp, typ, val, color in rows:
        swatch = f'<i class="sy-swatch {color}"></i>' if color else ""
        cells.append(f"<tr><td>{swatch}{_h(src)}</td><td>{_h(inp)}</td><td>{_h(typ)}</td><td>{_h(val)}</td></tr>")
    table = "" if not rows else ('<table class="sy-table"><tr><th>Source</th><th>Input</th><th>Type</th><th>Value</th></tr>'
                                 + "".join(cells) + "</table>")
    errs = "".join(f'<div class="sy-errors">{_h(e)}</div>' for e in g.get("errors", []))
    return (f'<div class="sy-report"><h2>System {_h(g.get("system", ""))}</h2>{errs}'
            f'<div class="sy-canvas">{system_svg(g)}</div>{table}{system_signals_html(g)}</div>')


def _test_files_state() -> dict:
    # {name: (mtime, size)} of the .test files in the working directory.
    state = {}
    try:
        for n in os.listdir(os.getcwd()):
            if n.endswith(".test") and os.path.isfile(n):
                st = os.stat(n)
                state[n] = (st.st_mtime_ns, st.st_size)
    except OSError:
        pass
    return state


def _written_card(before: dict):
    # The newest .test file changed since `before` that is a card, as .test JSON.
    after = _test_files_state()
    for n in sorted((n for n in after if before.get(n) != after[n]), key=lambda n: after[n][0], reverse=True):
        text = card_text(_read_text(n, _POINT_FILE_MAX))
        if text:
            return text
    return None


def _point_files_state(campaign_id: str, card_id: str, point_seq) -> dict:
    # {name: (mtime, size)} of the files named for this point, to see what a run wrote.
    pat = _point_file_pattern(campaign_id, card_id)
    state = {}
    try:
        for n in os.listdir(os.getcwd()):
            m = pat.match(n)
            if m and m.group(1) == str(point_seq) and os.path.isfile(n):
                st = os.stat(n)
                state[n] = (st.st_mtime_ns, st.st_size)
    except OSError:
        pass
    return state


def _ports_lookup(modules=None):
    """signet.check_wiring's ports(module): None if no registered module has that stem, else the ports its
    file declares ({} if none; text files under 1 MB only). A bad declaration raises ValueError."""
    names = {os.path.splitext(m)[0]: m for m in (modules if modules is not None else list_modules())}
    runnable = set(list_modules())

    def ports(stem):
        name = names.get(stem)
        if name is None or name not in runnable:
            return None
        path = os.path.join(_modules_dir(), name)
        try:
            if os.path.getsize(path) > 1_000_000:
                return {}
            with open(path, encoding="utf-8") as f:
                text = f.read()
        except (OSError, UnicodeDecodeError):
            return {}
        try:
            return signet.read_ports(text) or {}
        except ValueError as e:
            raise ValueError(f"{name}: {e}")
    return ports


def _wiring_errors(code, modules, live):
    """Raises ValueError listing the direction and port errors in a Signals box (see signet.check_wiring)."""
    if not signet or not str(code or "").strip():
        return
    plain = lambda t: os.path.isfile(t) or any(os.path.basename(str(x.get("file", ""))) == t for x in live or [])
    errs = signet.check_wiring(str(code), _ports_lookup(modules), plain)
    if errs:
        raise ValueError("\n".join(errs))


def _delivered(point_seq, card_id, campaign_id, live, points) -> dict:
    """{port: (point, input)} that other points' Signals boxes deliver to this point (connect(x, 20:volts),
    route exit), read now. A sequence's box delivers through its run(), so it is left out here."""
    get = getattr(signet, "deliveries", None)
    out = {}
    if not get or not isinstance(points, dict):
        return out
    me = str(point_seq)
    for seq, info in points.items():
        code = str((info or {}).get("signals") or "") if isinstance(info, dict) else ""
        if str(seq) == me or not code.strip() or signet.is_sequence(code):
            continue
        try:
            if not any(c["point"] in (me, "expr") for c in signet.connections(code, points=True) if not c["module"]):
                continue
        except ValueError:
            continue            # that point's own run reports its bad code
        try:
            host = _PointHost(card_id, campaign_id, live, _parse_params(info.get("params") or []), points)
            got = get(code, me, host)
        except ValueError as e:
            raise ValueError(f"point {seq}'s Signals: {e}")
        for port, v in got.items():
            if port in out:
                raise ValueError(f"{port} is set by both point {out[port][0]}'s Signals and point {seq}'s")
            out[port] = (str(seq), v)
    return out


def run_module(name: str, card_id: str = "", point_seq: str = "", campaign_id: str = "", args="", params=None,
               signals="", live=None, modules=None, held=None, points=None) -> dict:
    """Run one module from dist/ and wait for it. The name must be a bare
    filename that list_modules() reports -- never a path, so the browser can't
    point this at anything else. The module runs in the working directory
    (where .test files live) with STARFIRE_CAMPAIGN, STARFIRE_CARD and STARFIRE_POINT in its environment.
    The result's "files" lists the files named for this point that the run created or changed;
    "card" is a test card the run wrote (a .test file, else its whole output), as .test JSON.
    args is one command-line string, split into plain argv entries (no shell),
    so it reaches a module in any language the same way. params are the point's
    [[name, value]] rows, set as STARFIRE_PARAM_<NAME> (the page copies them into args).
    signals is the point's Signet code, run first against the campaign's live list
    ([{card, point, file}]); modules are the point's registered modules, for its checks.
    held is {port: {value, type, source}} a sequence holds for this module, merged with the point's own."""
    if name not in list_modules():
        return {"status": "failed", "error": f"{name!r} is not a runnable file in dist/"}
    try:
        extra = _parse_args(args)
    except ValueError as e:
        return {"status": "failed", "error": f"bad arguments: {e}"}
    try:
        param_items = _parse_params(params)
    except ValueError as e:
        return {"status": "failed", "error": f"bad parameters: {e}"}
    argv = [os.path.join(_modules_dir(), name)]
    in_signet = name.endswith(".signet")
    if in_signet and not signet:
        return {"status": "failed", "error": f"{_NO_SIGNET} (needed to run {name})"}
    interpreter = _INTERPRETERS.get(os.path.splitext(name)[1])
    if interpreter:
        found = shutil.which(interpreter)
        if not found:
            return {"status": "failed", "error": f"{interpreter} not found on PATH (needed to run {name})"}
        argv.insert(0, found)
    argv += extra
    stem = os.path.splitext(name)[0]
    if str(signals or "").strip() and not signet:
        return {"status": "failed", "error": f"{_NO_SIGNET}; clear the Signals box to run without it", "output": "", "files": []}
    host = _PointHost(card_id, campaign_id, live, param_items, points)
    host.point = str(point_seq)
    decl = None
    try:
        _wiring_errors(signals, modules, live)
        inputs = signet.evaluate(signals, stem, host, modules) if str(signals or "").strip() else {}
        for port in held or {}:
            if port in inputs:
                raise ValueError(f"{port} is set by both the sequence and this point's own Signals")
        inputs.update(held or {})
        for port, (src, v) in _delivered(point_seq, card_id, campaign_id, live, points).items():
            if port in inputs:
                raise ValueError(f"{port} is set by both point {src}'s Signals and this point's")
            inputs[port] = v
        if signet:
            decl = _ports_lookup(modules)(stem)
            errs = signet.contract_errors(stem, decl, inputs=inputs)
            if errs:
                raise ValueError("; ".join(errs))
    except (signet.Break if signet else ()) as e:
        return {"status": "broken", "error": str(e), "output": "".join(ln + "\n" for ln in host.lines), "files": []}
    except ValueError as e:
        return {"status": "failed", "error": f"Signet:\n{e}", "output": "".join(ln + "\n" for ln in host.lines), "files": []}
    inputs_path = f"{_point_prefix(campaign_id, card_id)}-{point_seq}-{stem}.inputs.json"
    env = _module_env(card_id, point_seq, campaign_id, param_items, stem,
                      os.path.join(os.getcwd(), inputs_path) if inputs else "")
    env["PYTHONUNBUFFERED"] = "1"   # so a Python module's output reaches its file as it prints
    out_path = _output_path(campaign_id, card_id, point_seq, name)
    key = (str(card_id), str(point_seq))
    started = time.monotonic()
    before = _point_files_state(campaign_id, card_id, point_seq)
    tests_before = _test_files_state()
    code_out, code_rc = [], 0
    if signet and os.path.splitext(name)[1].lower() in signet.SOURCE_TYPES:
        # C++, Rust, Ada or Fortran source: Signet builds it (once per change, cached) and gives the program to run
        try:
            argv = signet.build_module(argv[0], log=lambda t: host.lines.append("signet: " + t)) + extra
        except (LookupError, RuntimeError, OSError) as e:
            return {"status": "failed", "error": str(e), "output": "".join(ln + "\n" for ln in host.lines), "files": []}
    if in_signet:
        # A module written in Signet runs here, not in a process: it is bounded and holds no files open.
        try:
            with open(argv[0], encoding="utf-8") as f:
                signet.run_module_code(f.read(), inputs, f"{_point_prefix(campaign_id, card_id)}-{point_seq}-{stem}.",
                                       code_out.append)
        except signet.Break as e:
            code_rc = 1
            code_out.append(f"BREAK: {e}")
        except (ValueError, OSError) as e:
            code_rc = 1
            code_out.append(str(e))
    with _RUNS_LOCK:
        if key in _RUNS:
            return {"status": "failed", "error": "a module is already running on this point"}
        try:
            if signet:
                signet.write_inputs(inputs_path, inputs)
            out_file = open(out_path, "w", encoding="utf-8")   # the run's output, written as it goes
            out_file.write("".join(ln + "\n" for ln in host.lines + code_out))
            out_file.flush()
        except OSError as e:
            return {"status": "failed", "error": f"can't write {e.filename}: {e.strerror}"}
        if in_signet:
            proc = _Finished(code_rc)
        else:
            try:
                # Own session, so a timeout or BREAK can take down the module's children too.
                proc = subprocess.Popen(argv, cwd=os.getcwd(), env=env, stdout=out_file, stderr=subprocess.STDOUT,
                                        stdin=subprocess.DEVNULL, start_new_session=True)
            except OSError as e:
                out_file.close()
                return {"status": "failed", "error": f"could not start {name}: {e}"}
            _RUNS[key] = proc
    timed_out = False
    try:
        proc.wait(timeout=_MODULE_TIMEOUT_S)
    except subprocess.TimeoutExpired:
        timed_out = True
        os.killpg(proc.pid, signal.SIGKILL)
        proc.wait()
    finally:
        out_file.close()
        with _RUNS_LOCK:
            _RUNS.pop(key, None)
            broken = key in _BROKEN
            _BROKEN.discard(key)
    with open(out_path, encoding="utf-8", errors="replace") as f:
        out = f.read()
    if proc.returncode == 0 and not timed_out:
        status = "complete"
    else:
        status = "broken" if broken else "failed"
    result = {
        "status": status,
        "exit_code": proc.returncode,
        "seconds": round(time.monotonic() - started, 1),
        "output": _cap_lines(out),
        "files": sorted(n for n, st in _point_files_state(campaign_id, card_id, point_seq).items() if before.get(n) != st),
    }
    if status == "complete" and signet and decl:
        pre = f"{_point_prefix(campaign_id, card_id)}-{point_seq}-"
        errs = signet.contract_errors(stem, decl, written=[n[len(pre):] for n in result["files"] if n.startswith(pre) and n.endswith(".sig")])
        if errs:
            status = result["status"] = "failed"
            result["error"] = "; ".join(errs)
    if status == "complete" and signet and signet.STEM.match(stem):
        sig = signet.stamp(f"{_point_prefix(campaign_id, card_id)}-{point_seq}", stem, result["files"],
                           f"{campaign_id or 'XX-000'}:{card_id}-{point_seq}:")
        if any(sig.values()):
            result["signals"] = sig
            result["files"] = sorted(set(result["files"]) - set(sig["deleted"]))
    card = _written_card(tests_before) or card_text(out)
    if card:
        result["card"] = card
    if timed_out:
        result["error"] = f"timed out after {_MODULE_TIMEOUT_S} s"
    return result

class _Finished:
    """The stand-in for a process when a module ran inside Starfire (a .signet): already over."""

    def __init__(self, returncode):
        self.returncode = returncode

    def wait(self, timeout=None):
        return self.returncode


def _output_path(campaign_id, card_id, point_seq, name) -> str:
    # <prefix>-<point>-<module>.txt: a point file, so Signet code and reports can use it.
    stem = re.sub(r"[^A-Za-z0-9._-]", "_", os.path.splitext(str(name))[0]) or "module"
    return f"{_point_prefix(campaign_id, card_id)}-{point_seq}-{stem}.txt"


# Runs started from the page run in the background: (card, point) -> the latest job there.
# job["status"] is what the point shows; job["state"] is its module's own run. They differ while a sequence outlives it.
_JOBS: dict = {}
_JOB_IDS = iter(range(1, 1 << 62))
_SEQS: dict = {}      # (card, point) -> _SeqRunner, while that point's sequence is going
_WAITING: dict = {}   # (card, point) -> {module, inputs, seq, state}: a run a sequence asked for, held until RUN is pressed


def _wake(card_id):
    # a run on this card ended or RUN was pressed: its sequences step now
    with _RUNS_LOCK:
        runners = [r for k, r in _SEQS.items() if k[0] == str(card_id)]
    for r in runners:
        r.wake.set()


def _seq_line(ev, when) -> str:
    extra = []
    for k, v in ev.items():
        if k in ("t", "event", "line"):
            continue
        if k == "inputs":
            v = ", ".join(f"{p}={i['value']} ({i['type']})" for p, i in v.items()) or "none"
        extra.append(str(v) if k == "text" else f"{k}={v}")
    where = f"line {ev['line']}" if ev.get("line") else ""
    return f"{when:%Y-%m-%d %H:%M:%S.%f}"[:-3] + f"  {ev['t']:9.3f} s  {ev['event']:<11}  {where:<9} {'  '.join(extra)}".rstrip()


class _SeqHost(_PointHost):
    """What a sequence sees (Signet API 2): the card's points, and the run table."""

    def __init__(self, runner, campaign_id, live, params, pts):
        super().__init__(runner.key[0], campaign_id, live, params, pts)
        self.runner, self.campaign, self.live = runner, campaign_id, live
        self.started = set()

    def ask(self, prompt, type_name):
        answer = self.runner.ask("seq", prompt, type_name)
        self.runner.seq.event("input", None, text=f"{prompt} = {answer}")    # into the as-run log, in order
        return answer

    def echo(self, text):
        pass    # it is in the log already

    def resolve(self, text):
        try:
            return super().resolve(text)
        except LookupError:
            # a file still being written isn't ticked live: [card-]point:file is <prefix>-<point>-file
            m = re.match(r"(?:(.+)-)?(\d+):([A-Za-z0-9._-]+)\Z", text)
            if m and not m.group(3).endswith(".sig"):
                card, point, rest = m.group(1) or self.card, m.group(2), m.group(3)
                path = f"{_point_prefix(self.campaign, card)}-{point}-{rest}"
                if os.path.isfile(path):
                    return signet.source_value(path, f"{card}-{point}:{rest}", text)
            raise

    def run(self, point, inputs):
        point = str(point)
        info = self.pts.get(point)
        if not isinstance(info, dict):
            raise LookupError(f"no point {point} on this card")
        name = str(info.get("module") or "").strip()
        if not name:
            raise LookupError(f"point {point} has no module")
        if name not in (info.get("modules") or []) or name not in list_modules():
            raise LookupError(f"{name} isn't registered for point {point}")
        code, stem = str(info.get("signals") or ""), os.path.splitext(name)[0]
        if code.strip():
            if signet.is_sequence(code):
                raise RuntimeError(f"point {point} has a sequence of its own")
            try:
                own = {c["port"] for c in signet.connections(code) if c["module"] == stem}
            except ValueError as e:
                raise RuntimeError(f"point {point}'s Signals: {e}")
            for port in sorted(own & set(inputs)):
                raise RuntimeError(f"{port} is set by both this sequence and point {point}'s own Signals")
        key = (self.card, point)
        with _RUNS_LOCK:
            job = _JOBS.get(key)
            if key in _WAITING or (job and job["status"] == "running"):
                raise RuntimeError(f"point {point} is already running")
            _WAITING[key] = {"module": name, "inputs": inputs, "seq": self.runner, "state": "waiting"}   # until RUN
        self.started.add(point)

    def state(self, point):
        key = (self.card, str(point))
        with _RUNS_LOCK:
            w = _WAITING.get(key)
            if w and w["seq"] is self.runner:
                return w["state"]
            job = _JOBS.get(key)
            if job and job["started"] >= self.runner.t0:
                return job["state"]
        return "none"

    def break_runs(self):
        with _RUNS_LOCK:
            for k in [k for k, w in _WAITING.items() if w["seq"] is self.runner]:
                del _WAITING[k]
            ours = [p for p in self.started | {self.runner.key[1]}
                    if (_JOBS.get((self.card, p)) or {}).get("started", -1) >= self.runner.t0]
        for p in ours:
            _break_process(self.card, p)


class _SeqRunner:
    """One sequence going: Signet's Sequence, stepped by a thread about every second and whenever a run on
    the card ends or RUN is pressed on a waiting point. Its events are written to the as-run log."""

    def __init__(self, code, key, stem, campaign_id, live, params, pts):
        self.key, self.code, self.campaign = key, code, campaign_id
        self.wake, self.done, self.lock = threading.Event(), threading.Event(), threading.Lock()
        self.t0 = time.monotonic()
        self.host = _SeqHost(self, campaign_id, live, params, pts)
        self.seq = signet.Sequence(code, str(key[1]), self.host, module=stem, now=self.t0)   # ValueError: bad code
        self.points = self.seq.points
        self.status = signet.Status("waiting", "", "", None, [])
        self.log_name = f"{_point_prefix(campaign_id, key[0])}-{key[1]}-sequence.log"
        # input(): one question at a time, shown on the point's page. The point's own run reads the same top-level
        # lines as the sequence, so each asker takes answers in order and an answer given once serves both.
        self.question, self.answers, self.asked, self.stopping = None, [], {"run": 0, "seq": 0}, ""
        self.asking = threading.Condition()

    def begin(self):
        stamp = datetime.now(timezone.utc)
        head = [f"# Signet sequence on point {self.key[1]} of {self.key[0]} (campaign {self.campaign or 'XX-000'}), started {stamp:%Y-%m-%d %H:%M:%S} UTC",
                *("#   " + ln for ln in self.code.splitlines()), "#"]
        self._write(head, "w")

    def _write(self, lines, mode="a"):
        try:
            with open(self.log_name, mode, encoding="utf-8") as f:
                f.write("".join(ln + "\n" for ln in lines))
        except OSError:
            pass

    def record(self, st):
        with self.lock:
            if self.done.is_set():
                return
            when = datetime.now(timezone.utc)
            self._write([_seq_line(ev, when) for ev in st.events])
            self.status = st
            if st.state != "waiting":
                with _RUNS_LOCK:
                    if _SEQS.get(self.key) is self:
                        del _SEQS[self.key]
                    for k in [k for k, w in _WAITING.items() if w["seq"] is self]:
                        del _WAITING[k]
                self.done.set()

    def loop(self):
        while True:
            st = self.seq.step(time.monotonic())
            self.record(st)
            if st.state != "waiting":
                return
            self.wake.wait(1.0)
            self.wake.clear()

    def ask(self, who, prompt, type_name):
        with self.asking:
            n = self.asked[who]
            while True:
                if self.stopping:
                    raise LookupError(self.stopping)
                if n < len(self.answers) and self.answers[n][:2] == (prompt, type_name):
                    self.asked[who] = n + 1
                    return self.answers[n][2]
                if self.question is None:
                    n = len(self.answers)
                    self.question = {"prompt": prompt, "type": type_name}
                self.asking.wait(1.0)

    def answer(self, text) -> bool:
        with self.asking:
            q = self.question
            if q is None:
                return False
            self.answers.append((q["prompt"], q["type"], text))
            self.question = None
            self.asking.notify_all()
        self.wake.set()
        return True

    def stop(self, reason="broken by operator"):
        with self.asking:               # a wait on input() ends, so the step can finish
            self.stopping = reason
            self.asking.notify_all()
        self.record(self.seq.break_(time.monotonic(), reason))
        self.wake.set()

    def log_text(self, tail=60) -> str:
        try:
            with open(self.log_name, encoding="utf-8", errors="replace") as f:
                return "".join([ln for ln in f.readlines() if not ln.startswith("#")][-tail:])
        except OSError:
            return ""

    def finish(self, r) -> dict:
        """The module's result, with how the sequence ended; waits for the sequence to end."""
        self.done.wait()
        st, r = self.status, dict(r)
        r["sequence"] = {"state": st.state, "kind": st.kind, "reason": st.reason, "log": self.log_name}
        r["files"] = sorted(set(r.get("files") or []) | {self.log_name})
        if st.state == "broken":
            if r.get("status") == "complete":
                r["status"] = "broken"
            r["error"] = r.get("error") or st.reason
        r["output"] = (r.get("output") or "") + "\nSequence:\n" + self.log_text(200)
        return r


def stop_sequences():
    with _RUNS_LOCK:
        runners = list(_SEQS.values())
    for r in runners:
        r.stop("Starfire shut down")


def start_run(name, card_id="", point_seq="", campaign_id="", args="", params=None, signals="", live=None, modules=None,
              extra=None, points=None) -> dict:
    """Start run_module in the background and return at once: {"status": "running", "run": id}.
    run_status() reports progress and, once it ends, the result. A point whose Signals box is a sequence
    also starts the sequence (points: every point of the card, for its run() and point parameters).
    `extra` is the connections a sequence held for the module, given when RUN is pressed on it."""
    key = (str(card_id), str(point_seq))
    code, runner = str(signals or ""), None
    fail = lambda msg: {"status": "failed", "error": msg, "output": "", "files": []}
    if signet and code.strip() and key not in _WAITING and signet.is_sequence(code):
        if name not in list_modules():
            return fail(f"{name!r} is not a runnable file in dist/")
        try:
            _wiring_errors(code, modules, live)
            runner = _SeqRunner(code, key, os.path.splitext(str(name))[0], campaign_id, live,
                                _parse_params(params), points if isinstance(points, dict) else {})
        except ValueError as e:
            return fail(f"Signet:\n{e}")
    with _RUNS_LOCK:
        job = _JOBS.get(key)
        if job and job["status"] == "running":
            return fail("a module is already running on this point")
        held = _WAITING.get(key)    # a run a sequence asked for: this RUN starts it
        if held and held["state"] != "waiting":
            return fail(f"this run was {held['state']}")
        if not held:
            for k, r in _SEQS.items():
                if k[0] == key[0] and key[1] in r.points:
                    return fail(f"point {k[1]}'s sequence uses this point")
        if runner:
            for k, r in _SEQS.items():
                if k[0] == key[0] and set(r.points) & set(runner.points):
                    return fail(f"point {k[1]}'s sequence is going and uses points {', '.join(sorted(set(r.points) & set(runner.points)))}")
            for p in runner.points:
                busy = _JOBS.get((key[0], p))
                if p != key[1] and (busy and busy["status"] == "running" or (key[0], p) in _WAITING):
                    return fail(f"point {p} is running, and the sequence uses it")
            _SEQS[key] = runner
        if held:
            extra = held["inputs"]
            del _WAITING[key]
        job = {"run": next(_JOB_IDS), "status": "running", "state": "running", "module": str(name), "started": time.monotonic(),
               "output_file": _output_path(campaign_id, card_id, point_seq, name), "result": None, "runner": runner}
        _JOBS[key] = job
    if held:
        _wake(card_id)

    def work():
        try:
            r = run_module(name, card_id, point_seq, campaign_id, args, params, signals, live, modules, extra, points)
        except Exception as e:   # keep the job from hanging in "running"
            r = {"status": "failed", "error": f"internal error: {e}", "output": "", "files": []}
        state = r.get("status") or "failed"
        with _RUNS_LOCK:
            job["state"] = state
            if not runner:
                job.update(status=state, result=r)
        _wake(card_id)
        if runner:
            r = runner.finish(r)
            with _RUNS_LOCK:
                job.update(status=r.get("status") or "failed", result=r)
    threading.Thread(target=work, daemon=True).start()
    if runner:
        runner.begin()
        threading.Thread(target=runner.loop, daemon=True).start()
    return {"status": "running", "run": job["run"]}


def run_status(card_id, point_seq=None, tail: int = 200) -> dict:
    """One point's latest run: {run, status, module, seconds, output so far | result}; or, with no point,
    {"points": {point: {run, status, module}}} for every point of the card that has run since starfire started.
    A run a sequence is holding for RUN is "waiting"; a point with a sequence going also has "sequence"."""
    card = str(card_id)
    with _RUNS_LOCK:
        jobs = {k: dict(v) for k, v in _JOBS.items() if k[0] == card}
        waits = {k[1]: dict(w) for k, w in _WAITING.items() if k[0] == card}
        seqs = {k[1]: r for k, r in _SEQS.items() if k[0] == card}
    seq_info = lambda r: {"state": r.status.state, "reason": r.status.reason, "kind": r.status.kind, "log": r.log_text(40),
                          "question": r.question}
    if point_seq is None:
        pts = {k[1]: {"run": j["run"], "status": j["status"], "module": j["module"]} for k, j in jobs.items()}
        for p, w in waits.items():
            pts[p] = {"status": w["state"], "module": w["module"]}
        for p, r in seqs.items():
            pts.setdefault(p, {})["sequence"] = {"state": r.status.state, "reason": r.status.reason}
        return {"points": pts}
    p = str(point_seq)
    j = jobs.get((card, p))
    if p in waits and not (j and j["status"] == "running"):
        return {"status": waits[p]["state"], "module": waits[p]["module"], "seconds": 0}
    if not j:
        return {"status": "none"}
    out = {"run": j["run"], "status": j["status"], "module": j["module"],
           "seconds": round(time.monotonic() - j["started"], 1)}
    if j["status"] == "running":
        try:
            with open(j["output_file"], encoding="utf-8", errors="replace") as f:
                out["output"] = "".join(f.readlines()[-tail:])
        except OSError:
            out["output"] = ""
        if p in seqs:
            out["sequence"] = seq_info(seqs[p])
        elif j.get("runner"):
            out["sequence"] = seq_info(j["runner"])
    else:
        out["result"] = j["result"]
    return out


def _break_process(card_id, point_seq, grace: float = 3.0) -> dict:
    # SIGTERM the module running on this point; SIGKILL if it is still alive after `grace` s.
    key = (str(card_id), str(point_seq))
    with _RUNS_LOCK:
        proc = _RUNS.get(key)
        if proc is None:
            return {"success": False, "error": "no module is running on this point"}
        _BROKEN.add(key)
    try:
        os.killpg(proc.pid, signal.SIGTERM)
    except ProcessLookupError:
        return {"success": True}

    def _kill():
        if proc.poll() is None:
            try:
                os.killpg(proc.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
    threading.Timer(grace, _kill).start()
    return {"success": True}


def break_module(card_id, point_seq, grace: float = 3.0) -> dict:
    """BREAK on a point: stops its sequence (and the runs it started), or the run a sequence is holding
    for RUN (its sequence breaks on its next step), or the module running there."""
    key = (str(card_id), str(point_seq))
    with _RUNS_LOCK:
        runner, held = _SEQS.get(key), _WAITING.get(key)
        if held and not runner and key not in _RUNS:
            held["state"] = "broken"
    if runner:
        runner.stop("broken by operator")
        return {"success": True}
    if held and held["state"] == "broken":
        held["seq"].wake.set()
        return {"success": True}
    return _break_process(card_id, point_seq, grace)


def _fill_flags(html: str) -> str:
    return (html.replace("__JIRA_AVAILABLE__", "true" if JIRA_AVAILABLE else "false")
                .replace("__SIGNET_AVAILABLE__", "true" if SIGNET_AVAILABLE else "false"))

_TEST_HTML = _fill_flags(_read_static("test.html"))
_CAMPAIGN_HTML = _fill_flags(_read_static("campaign.html"))
_SYSTEM_HTML = _read_static("system.html")

# Calculator: math and metrics on a point's csv columns, in validate.py's expression syntax
# (@voltage, max @voltage, @voltage * @current) but no limits. Needs numpy, loaded on first use.
_CALC_FUNCS = ("max", "min", "mean", "median", "sd", "rms", "range", "abs")
_CALC_TOKEN = re.compile(r"\s*(?:(\d+\.?\d*(?:[eE][+-]?\d+)?|\.\d+(?:[eE][+-]?\d+)?)|@([A-Za-z_]\w*)|([()+\-*/])|([A-Za-z_]\w*))")


class CalcError(Exception):
    pass


def _calc_resolve(name: str, headers: list) -> str:
    n = name.lower()
    hits = [h for h in headers if h.lower() == n] or [h for h in headers if h.lower().split("_")[0] == n]
    if len(hits) == 1:
        return hits[0]
    raise CalcError(f"@{name} matches {', '.join(hits)}" if hits else f"no column @{name}")


def _calc_tokens(expr: str) -> list:
    toks, pos = [], 0
    while expr[pos:].strip():
        m = _CALC_TOKEN.match(expr, pos)
        if not m:
            c = expr[pos:].strip()[0]
            raise CalcError("limits (< > == AND OR NOT) belong in validate.py" if c in "<>=!" else f"unexpected '{c}'")
        num, ref, op, word = m.groups()
        if word and word.lower() not in _CALC_FUNCS:
            raise CalcError("limits (< > == AND OR NOT) belong in validate.py" if word.lower() in ("and", "or", "not", "any", "all")
                            else f"unknown word '{word}'")
        toks.append(("num", float(num)) if num else ("ref", ref) if ref else ("op", op) if op else ("fn", word.lower()))
        pos = m.end()
    return toks


def _calc_eval(expr: str, table: dict, np):
    # Recursive descent: + - > * / > unary minus, prefix function > atom. Prefix functions bind tightest.
    toks, headers, i = _calc_tokens(expr), list(table), [0]
    if not toks:
        raise CalcError("empty")

    def take(*vals):
        t = toks[i[0]] if i[0] < len(toks) else None
        if t and t[0] in ("op", "fn") and (not vals or t[1] in vals):
            i[0] += 1
            return t
        return None

    def add():
        v = mul()
        while (t := take("+", "-")):
            v = v + mul() if t[1] == "+" else v - mul()
        return v

    def mul():
        v = unary()
        while (t := take("*", "/")):
            r = unary()
            if t[1] == "/" and np.any(r == 0):
                raise CalcError("division by zero")
            v = v * r if t[1] == "*" else v / r
        return v

    def unary():
        if take("-"):
            return -unary()
        t = take(*_CALC_FUNCS)
        if t:
            v = unary()
            f = {"max": np.max, "min": np.min, "mean": np.mean, "median": np.median, "sd": np.std, "range": np.ptp,
                 "rms": lambda x: np.sqrt(np.mean(np.square(x))), "abs": np.abs}[t[1]]
            return np.asarray(f(v))
        return atom()

    def atom():
        if i[0] >= len(toks):
            raise CalcError("expression ends early")
        t = toks[i[0]]
        i[0] += 1
        if t[0] == "num":
            return np.asarray(t[1])
        if t[0] == "ref":
            return table[_calc_resolve(t[1], headers)]
        if t[1] == "(":
            v = add()
            if not take(")"):
                raise CalcError("missing ')'")
            return v
        raise CalcError(f"unexpected '{t[1]}'")

    with np.errstate(all="ignore"):
        v = add()
    if i[0] < len(toks):
        raise CalcError(f"unexpected '{toks[i[0]][1]}'")
    return v


def _calc_table(path: str, np):
    # {header: array} of the all-numeric columns of a csv or csv-like .txt (first row is the header
    # unless it is all numbers; "#" lines ignored), read in full.
    def isnum(s):
        try:
            return np.isfinite(float(s))
        except ValueError:
            return False
    with open(path, encoding="utf-8", errors="replace") as f:
        rows = [[c.strip() for c in ln.rstrip("\r\n").split(",")] for ln in f if ln.strip() and not ln.lstrip().startswith("#")]
    if not rows:
        return {}
    head = rows[0] if not all(isnum(c) for c in rows[0]) else None
    data = rows[1:] if head else rows
    out = {}
    for i in range(max(len(r) for r in rows)):
        h = head[i] if head and i < len(head) and head[i] else f"col{i + 1}"
        cells = [r[i] if i < len(r) else "" for r in data]
        if cells and all(isnum(c) for c in cells):
            out[h] = np.array([float(c) for c in cells])
    return out


def calculate(campaign_id: str, card_id: str, name, exprs: list) -> dict:
    """{"columns": [...], "results": [{"expr", "value"} or {"expr", "error"}]} for a point's csv file.
    An array result shows its size and first values; a scalar shows 6 significant figures."""
    try:
        import numpy as np
    except ImportError:
        return {"error": "The calculator needs numpy: pip install numpy"}
    if not _own_file(campaign_id, card_id, name):
        return {"error": "file not found"}
    try:
        table = _calc_table(os.path.join(os.getcwd(), name), np)
    except OSError as e:
        return {"error": str(e)}
    if not table:
        return {"error": f"{name} has no numeric columns"}
    results = []
    for e in exprs if isinstance(exprs, list) else []:
        e = str(e)
        if not e.strip():
            continue
        try:
            v = _calc_eval(e, table, np)
            vals = np.atleast_1d(v)
            g = lambda x: f"{float(x):.6g}"
            results.append({"expr": e, "value": g(v) if v.ndim == 0
                            else f"{vals.size} samples: " + ", ".join(g(x) for x in vals[:8]) + (" ..." if vals.size > 8 else "")})
        except CalcError as err:
            results.append({"expr": e, "error": str(err)})
    return {"columns": list(table), "results": results}


_POINT_FILE_EXTS = ("csv", "txt", "test")
_POINT_FILE_MAX = 5_000_000   # characters

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class _Handler(BaseHTTPRequestHandler):
    def log_message(self, format, *args):
        pass

    def _send(self, code, ctype, body: bytes):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _send_json(self, code, obj):
        self._send(code, "application/json", json.dumps(obj).encode())

    def _body(self) -> bytes:
        return self.rfile.read(int(self.headers.get("Content-Length", 0)))

    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path
        if path == "/":
            self._send(200, "text/html", _CAMPAIGN_HTML.encode())
        elif path == "/test":
            # ?load=<file> opens a specific .test file (read client-side via /load).
            self._send(200, "text/html", _TEST_HTML.encode())
        elif path == "/load":
            # Read a file by name from the working directory -- how a
            # test tab or the campaign page opens a saved file.
            qs = parse_qs(parsed.query)
            safe = os.path.basename((qs.get("file") or [""])[0])
            file_path = os.path.join(os.getcwd(), safe)
            if not safe or not os.path.isfile(file_path):
                self._send_json(404, {"error": f"not found: {safe!r}"})
                return
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    self._send(200, "application/json", f.read().encode())
            except OSError as e:
                self._send_json(500, {"error": str(e)})
        elif path == "/system":
            # ?card=<id> draws that card's System; the card and live list come from the opening page.
            self._send(200, "text/html", _SYSTEM_HTML.encode())
        elif path == "/run_status":
            # ?card=&point= for one point's latest run, or ?card= alone for every point of the card
            qs = parse_qs(parsed.query)
            point = (qs.get("point") or [None])[0]
            self._send_json(200, run_status((qs.get("card") or [""])[0], point))
        elif path == "/signal_value":
            qs = parse_qs(parsed.query)
            self._send_json(200, signal_value((qs.get("file") or [""])[0]))
        elif path == "/list_modules":
            self._send_json(200, {"modules": list_modules()})
        elif path == "/point_files":
            qs = parse_qs(parsed.query)
            try:
                attached = json.loads(qs["attached"][0]) if "attached" in qs else None
            except json.JSONDecodeError:
                attached = {}
            self._send_json(200, {"files": list_point_files((qs.get("campaign") or [""])[0], (qs.get("card") or [""])[0],
                                                              attached if isinstance(attached, dict) or attached is None else {})})
        elif path == "/list_files":
            # Lists files in the working directory by extension.
            qs = parse_qs(parsed.query)
            ext = (qs.get("ext") or [""])[0]
            if not re.match(r"^\.[A-Za-z0-9]+$", ext or ""):
                self._send_json(400, {"error": "ext must look like '.test'"})
                return
            try:
                names = sorted(
                    f for f in os.listdir(os.getcwd())
                    if f.endswith(ext) and os.path.isfile(f)
                )
            except OSError as e:
                self._send_json(500, {"error": str(e)})
                return
            self._send_json(200, {"files": names})
        elif path.startswith("/static/"):
            self._serve_static(path[len("/static/"):])
        else:
            self._send(404, "text/plain", b"not found")

    def _serve_static(self, filename: str) -> None:
        safe = os.path.basename(filename)
        file_path = os.path.join(_static_dir(), safe)
        if not safe or not os.path.isfile(file_path):
            self._send(404, "text/plain", b"not found")
            return
        content_type = "text/css" if safe.endswith(".css") else "application/octet-stream"
        try:
            with open(file_path, "rb") as f:
                self._send(200, content_type, f.read())
        except OSError:
            self._send(500, "text/plain", b"error reading static file")

    def _write_file(self, filename: str, obj) -> None:
        # Saved to the server's working directory, meant to be re-opened by this program.
        out_path = os.path.join(os.getcwd(), filename)
        try:
            with open(out_path, "w", encoding="utf-8") as f:
                json.dump(obj, f, indent=2)
        except OSError as e:
            self._send_json(500, {"saved": False, "error": str(e)})
            return
        self._send_json(200, {"saved": True, "filename": filename, "path": out_path})

    def _handle_capture_image(self, raw_body: bytes):
        # Payload is {"data": "data:image/jpeg;base64,..."}, not a card.
        try:
            data_url = json.loads(raw_body)["data"]
        except (json.JSONDecodeError, KeyError, TypeError) as e:
            self._send_json(400, {"error": f"bad request: {e}"})
            return
        try:
            _header, b64data = data_url.split(",", 1)
            raw = base64.b64decode(b64data)
        except (ValueError, binascii.Error) as e:
            self._send_json(400, {"error": f"bad image data: {e}"})
            return
        # Millisecond resolution to deconflict a quick burst of captures.
        ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:-3]
        filename = f"{ts}.jpg"
        out_path = os.path.join(os.getcwd(), filename)
        try:
            with open(out_path, "wb") as f:
                f.write(raw)
        except OSError as e:
            self._send_json(500, {"saved": False, "error": str(e)})
            return
        self._send_json(200, {"saved": True, "filename": filename, "path": out_path})

    def _handle_create_jira_ticket(self, raw_body: bytes):
        # {"summary", "description", "epic_key", "issue_type"}; issue_type
        # optional -- falls back to cfg["issue_type"] if omitted.
        if not JIRA_AVAILABLE:
            self._send_json(404, {"success": False, "error": "Jira integration not available"})
            return
        try:
            payload = json.loads(raw_body)
        except json.JSONDecodeError as e:
            self._send_json(400, {"success": False, "error": f"bad request: {e}"})
            return
        try:
            result = jira_create_ticket(
                payload.get("summary", ""), payload.get("description", ""),
                payload.get("epic_key", ""), payload.get("issue_type") or None,
            )
        except Exception as e:
            result = {"success": False, "error": f"Jira ticket creation raised: {e}"}
        self._send_json(200, result)

    def _local_json_payload(self, raw_body: bytes):
        """For routes that run or write: only accept a request a web page on
        another site couldn't have sent, a JSON content type (which forces a CORS
        preflight this server never answers) and our own Host header (which stops
        DNS rebinding). Returns the payload dict, or None after sending the error."""
        port = self.server.server_address[1]
        if (self.headers.get("Content-Type", "").split(";")[0].strip() != "application/json"
                or self.headers.get("Host", "") not in (f"localhost:{port}", f"127.0.0.1:{port}")):
            self._send_json(403, {"status": "failed", "error": "forbidden"})
            return None
        try:
            payload = json.loads(raw_body)
        except json.JSONDecodeError as e:
            self._send_json(400, {"status": "failed", "error": f"bad request: {e}"})
            return None
        if not isinstance(payload, dict):
            self._send_json(400, {"status": "failed", "error": "bad request: expected a JSON object"})
            return None
        return payload

    def _handle_run_module(self, raw_body: bytes):
        payload = self._local_json_payload(raw_body)
        if payload is None:
            return
        self._send_json(200, start_run(str(payload.get("module") or ""), payload.get("card", ""),
                                        payload.get("point", ""), payload.get("campaign", ""),
                                        payload.get("args", ""), payload.get("params"),
                                        payload.get("signals", ""), payload.get("live"), payload.get("modules"),
                                        points=payload.get("points")))

    def _handle_add_card_file(self, raw_body: bytes):
        # {"campaign", "card", "point", "name", "data": base64} -> <campaign>-<card>-<point>-<name>
        payload = self._local_json_payload(raw_body)
        if payload is None:
            return
        try:
            data = base64.b64decode(payload.get("data") or "", validate=True)
        except (binascii.Error, ValueError, TypeError):
            self._send_json(400, {"saved": False, "error": "bad request"})
            return
        r = add_card_file(str(payload.get("campaign") or ""), str(payload.get("card") or ""),
                          payload.get("point"), payload.get("name"), data)
        self._send_json(200 if r["saved"] else 400, r)

    def _handle_break_module(self, raw_body: bytes):
        payload = self._local_json_payload(raw_body)
        if payload is None:
            return
        self._send_json(200, break_module(payload.get("card", ""), payload.get("point", "")))

    def _handle_calc(self, raw_body: bytes):
        # {"campaign", "card", "file", "exprs": [...]} -> {"columns", "results"} (see calculate)
        payload = self._local_json_payload(raw_body)
        if payload is None:
            return
        self._send_json(200, calculate(str(payload.get("campaign") or ""), str(payload.get("card") or ""),
                                       payload.get("file"), payload.get("exprs")))

    def _handle_save_point_file(self, raw_body: bytes):
        # {"campaign", "card", "point", "ext": "csv"|"txt"|"test", "content"} -> <campaign>-<card>-<point>.<ext>
        payload = self._local_json_payload(raw_body)
        if payload is None:
            return
        ext = payload.get("ext")
        content = payload.get("content")
        if ext not in _POINT_FILE_EXTS or not isinstance(content, str) or len(content) > _POINT_FILE_MAX:
            self._send_json(400, {"saved": False, "error": "bad request"})
            return
        if ext == "test" and card_text(content) is None:
            self._send_json(400, {"saved": False, "error": "not a test card"})
            return
        stem = "-".join(re.sub(r"[^A-Za-z0-9._-]", "_", str(payload.get(k) or d))
                        for k, d in (("campaign", "XX-000"), ("card", "TC"), ("point", "0")))
        if ext == "txt" and content.split("\n", 1)[0].strip() == "# validate":
            stem += "-VAL"    # validate.py output gets its own file; the point's module-output .txt stays
        filename = _safe_filename(stem, "." + ext)
        try:
            with open(os.path.join(os.getcwd(), filename), "w", encoding="utf-8", newline="") as f:
                f.write(content)
        except OSError as e:
            self._send_json(500, {"saved": False, "error": str(e)})
            return
        self._send_json(200, {"saved": True, "filename": filename})

    def do_POST(self):
        path = urlparse(self.path).path
        raw = self._body()
        if path == "/capture_image":
            return self._handle_capture_image(raw)
        if path == "/create_jira_ticket":
            return self._handle_create_jira_ticket(raw)
        if path == "/run_module":
            return self._handle_run_module(raw)
        if path == "/break_module":
            return self._handle_break_module(raw)
        if path == "/seq_answer":
            return self._handle_seq_answer(raw)
        if path == "/save_point_file":
            return self._handle_save_point_file(raw)
        if path == "/calc":
            return self._handle_calc(raw)
        if path == "/add_card_file":
            return self._handle_add_card_file(raw)
        try:
            data = json.loads(raw)
        except json.JSONDecodeError as e:
            self._send_json(400, {"error": str(e)})
            return
        try:
            self._route_json(path, data)
        except (TypeError, KeyError, AttributeError, ValueError) as e:
            self._send_json(400, {"error": str(e)})

    def _handle_seq_answer(self, raw) -> None:
        # input() in a sequence: {"card", "point", "answer"} for the sequence asking on that point
        try:
            data = json.loads(raw)
            key = (str(data["card"]), str(data["point"]))
            answer = str(data["answer"])
        except (ValueError, KeyError, TypeError) as e:
            return self._send_json(400, {"error": f"bad answer: {e}"})
        with _RUNS_LOCK:
            r = _SEQS.get(key)
        if not r or not r.answer(answer):
            return self._send_json(409, {"error": "no question is waiting on this point"})
        self._send_json(200, {"ok": True})

    def _route_json(self, path: str, data) -> None:
        host = self.headers.get("Host", "localhost")
        # Test-card routes
        if path in ("/export/csv", "/export/print", "/export/test/docx", "/save"):
            card = TestCard.from_dict(data)
            if path == "/export/csv":
                self._send(200, "text/csv", export_csv(card).encode())
            elif path == "/export/print":
                self._send(200, "text/html", export_print_html(procedure_only(card), host).encode())
            elif path == "/export/test/docx":
                self._send(200, _DOCX_MIME, export_test_docx(procedure_only(card)))
            else:
                self._write_file(_safe_filename(card.id, ".test"), card.to_dict())
        # Campaign routes
        elif path in ("/save_campaign", "/export/campaign/print", "/export/campaign/docx", "/vcrm", "/campaign_status"):
            camp = Campaign.from_dict(data)
            if path == "/save_campaign":
                self._write_file(_safe_filename(camp.id, ".camp"), camp.to_dict())
            elif path == "/export/campaign/print":
                self._send(200, "text/html", export_campaign_print_html(camp, host).encode())
            elif path == "/export/campaign/docx":
                self._send(200, _DOCX_MIME, export_campaign_docx(camp))
            elif path == "/campaign_status":
                self._send_json(200, campaign_statuses(camp))
            else:
                self._send_json(200, compute_vcrm(camp))
        # System graph: {"card": <the .test dict>, "campaign", "live": [{card, point, file}]}
        # Measure column: {"card", "campaign", "point", "live"}
        elif path == "/point_signals":
            self._send_json(200, point_signals(data.get("card") or {}, str(data.get("campaign") or ""),
                                               data.get("point"), data.get("live") or []))
        # Connection checkboxes: {"points": [{seq, module, signals}]}
        elif path == "/port_map":
            self._send_json(200, port_map(data.get("points") or []))
        elif path == "/system_graph":
            g = system_graph(data.get("card") or {}, str(data.get("campaign") or ""), data.get("live") or [])
            self._send_json(200, dict(g, svg=system_svg(g), signals=system_signals_html(g)))
        # A saved .sys redrawn: the System graph as exported (no svg, no signals)
        elif path == "/system_view":
            if (data.get("format") != "starfire-system" or not isinstance(data.get("boxes"), list)
                    or not isinstance(data.get("lines"), list)):
                self._send_json(400, {"error": "not a Starfire .sys file"})
            else:
                g = dict(data, errors=[str(e) for e in data.get("errors") or []])
                try:
                    self._send_json(200, dict(g, svg=system_svg(g), signals=""))
                except (KeyError, TypeError, ValueError, AttributeError) as e:
                    self._send_json(400, {"error": f"bad .sys file ({type(e).__name__}: {e})"})
        # BRANCH / EXTEND: {"mode", "card", "point", "file", "campaign", "used"}
        elif path == "/feed_card":
            self._send_json(200, feed_card(data))
        # Create test: {"campaign": <the .camp dict>, "requirements": [codes]}
        elif path == "/create_test":
            self._send_json(200, create_test(Campaign.from_dict(data["campaign"]), data["requirements"]))
        # Test Report routes: {"card", "campaign", "files": {seq: [names]}, "preview"}
        elif path in ("/export/report/print", "/export/report/docx"):
            card = TestCard.from_dict(data["card"])
            camp_id, files = str(data.get("campaign") or ""), data.get("files") or {}
            if path == "/export/report/print":
                html = export_report_html(card, camp_id, files, host, preview=bool(data.get("preview")), live=data.get("live"))
                self._send(200, "text/html", html.encode())
            else:
                self._send(200, _DOCX_MIME, export_report_docx(card, camp_id, files))
        # Campaign report (SR) routes: {"campaign": <the .camp dict>, "preview"}
        elif path in ("/export/campaign_report/print", "/export/campaign_report/docx"):
            camp = Campaign.from_dict(data["campaign"])
            if path.endswith("/print"):
                self._send(200, "text/html", export_campaign_report_html(camp, host, bool(data.get("preview"))).encode())
            else:
                self._send(200, _DOCX_MIME, export_campaign_report_docx(camp))
        # VCRM routes (the VCRM is plain JSON, not a dataclass)
        elif path == "/save_vcrm":
            self._write_file(vcrm_filename(data.get("campaign_id", "VCRM")), data)
        elif path == "/export/vcrm/print":
            self._send(200, "text/html", export_vcrm_print_html(data, host).encode())
        elif path == "/export/vcrm/csv":
            self._send(200, "text/csv", export_vcrm_csv(data).encode())
        elif path == "/export/vcrm/docx":
            self._send(200, _DOCX_MIME, export_vcrm_docx(data))
        else:
            self._send(404, "text/plain", b"not found")


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--port", type=int, default=5791)
    p.add_argument("--no-browser", action="store_true")
    args = p.parse_args()

    if hasattr(os, "geteuid") and os.geteuid() == 0:
        print("WARNING: Running as root. sudo is not needed to run this --")
        print("         only to clear out a root-owned dist/ from an earlier")
        print("         build, if that ever comes up. Run it as your normal")
        print("         user instead: root has no desktop session of its own,")
        print("         so the browser likely will fail to open correctly,")
        print("         and any files this writes would end up root-owned as well.\n")

    url = f"http://localhost:{args.port}"
    try:
        server = ThreadingHTTPServer(("localhost", args.port), _Handler)
    except OSError as e:
        print(f"ERROR: can't listen on port {args.port}: {e.strerror}.")
        if e.errno == errno.EADDRINUSE:
            print(f"       Starfire (or something else) is already running there: open {url},")
            print("       stop it first, or use --port to pick another.")
        return 1
    print(f"Serving GUI at {url}")
    print("Press Ctrl+C to stop\n")
    if not args.no_browser:
        threading.Timer(0.4, lambda: webbrowser.open(url)).start()

    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print()
        print("Shutting down")
    stop_sequences()
    return 0


if __name__ == "__main__":
    sys.exit(main())
